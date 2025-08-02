import os
import time
import tempfile
from typing import Dict, Any, List
import docx
from docx import Document
from collections import OrderedDict

# Define a constant for the character limit for easy changes
SLIDE_CHAR_LIMIT = 600

# ==============================================================================
# 1. NEW, SMARTER HELPER FUNCTIONS
# ==============================================================================

def group_paragraphs_by_heading(doc_path: str) -> OrderedDict:
    """
    Reads a .docx file and groups paragraphs under their respective headings.
    This preserves the document's structure.

    Returns:
        An OrderedDict mapping heading_text -> [paragraph_text_1, paragraph_text_2, ...]
    """
    print("Step 1: Grouping content by document headings...")
    document = Document(doc_path)
    grouped_content = OrderedDict()
    current_heading = "Introduction" # Default heading for content before the first real heading
    
    # Ensure the default heading exists
    grouped_content[current_heading] = []

    for p in document.paragraphs:
        # Check if the paragraph is a heading
        if p.style.name.startswith('Heading'):
            current_heading = p.text.strip()
            if current_heading not in grouped_content:
                grouped_content[current_heading] = []
        # Add non-empty paragraph text to the current heading's list
        elif p.text.strip():
            grouped_content[current_heading].append(p.text.strip())

    # Clean up any headings that ended up with no content
    return OrderedDict([(h, p) for h, p in grouped_content.items() if p])


def split_long_section_into_slides(heading: str, paragraphs: List[str]) -> List[str]:
    """
    Takes a single, oversized section and splits it into multiple slides,
    each with the same heading.
    """
    print(f"    -> Section '{heading}' is too long. Splitting structurally...")
    
    sub_slides = []
    current_slide_content = ""
    heading_len = len(heading) + 2 # Account for heading and two newlines

    for p in paragraphs:
        # If the current slide is empty and this single paragraph is too long,
        # it has to go on its own slide. (Edge Case)
        if not current_slide_content and (heading_len + len(p) > SLIDE_CHAR_LIMIT):
            print(f"       - Paragraph starting with '{p[:30]}...' is too long, creating a dedicated slide.")
            sub_slides.append(f"{heading}\n\n{p}")
            continue

        # If adding the next paragraph exceeds the limit, finalize the current slide
        if current_slide_content and (heading_len + len(current_slide_content) + len(p) + 2 > SLIDE_CHAR_LIMIT):
            print(f"       - Finalizing a sub-slide (length: {heading_len + len(current_slide_content)}).")
            sub_slides.append(f"{heading}\n\n{current_slide_content}")
            current_slide_content = p # Start the new slide with the current paragraph
        else:
            # Otherwise, add the paragraph to the current slide
            if not current_slide_content:
                current_slide_content = p
            else:
                current_slide_content += f"\n\n{p}"

    # Add the last remaining slide content
    if current_slide_content:
        print(f"       - Finalizing the last sub-slide (length: {heading_len + len(current_slide_content)}).")
        sub_slides.append(f"{heading}\n\n{current_slide_content}")

    return sub_slides


def create_smarter_slide_list(grouped_content: OrderedDict) -> List[str]:
    """
    Creates a final list of slides from content grouped by headings.
    It intelligently decides whether to keep a section on one slide or split it.
    """
    print("\nStep 2: Creating final slide list from grouped content...")
    final_slides = []
    
    for heading, paragraphs in grouped_content.items():
        # Combine all paragraphs into one block to check total length
        full_content_text = "\n\n".join(paragraphs)
        total_len = len(heading) + len(full_content_text) + 2

        print(f"  - Processing section '{heading}' (total length: {total_len})")

        # Case 1: The entire section fits on one slide
        if total_len <= SLIDE_CHAR_LIMIT:
            print("    -> Section fits on a single slide.")
            slide = f"{heading}\n\n{full_content_text}"
            final_slides.append(slide)
        # Case 2: The section is too long and must be split
        else:
            split_slides = split_long_section_into_slides(heading, paragraphs)
            final_slides.extend(split_slides)
            
    return final_slides


# ==============================================================================
# 2. YOUR MAIN GRAPH NODE (Updated to use new functions)
# ==============================================================================

def initialize_and_prepare(state: Dict[str, Any]) -> Dict[str, Any]:
    """
    Node khởi tạo: Đọc DOCX, chia chunks thông minh, và thiết lập trạng thái.
    """
    print("\n---------------- Executing Node: initialize_and_prepare ----------------")
    docx_file = state['docx_file']
    
    # Use the new, smarter functions
    grouped_content = group_paragraphs_by_heading(docx_file)
    slide_list = create_smarter_slide_list(grouped_content)
    
    print(f"\nFinal Result: Đã tạo {len(slide_list)} slide chunks ban đầu.")
    
    temp_dir = tempfile.mkdtemp()

    return {
        "slide_chunks": slide_list,
        "current_slide_index": -1,
        "final_slides": [],
        "all_debug_info": [],
        "previous_slide_context": None,
        "temp_dir": temp_dir,
    }

# ==============================================================================
# 3. MAIN EXECUTION SCRIPT
# ==============================================================================

if __name__ == "__main__":
    # Create a more challenging sample docx file
    sample_docx_path = "nature.docx"
    doc = Document()
    
    # Section 1: Short, should be one slide
    doc.add_heading('Slide 1: Giới thiệu về LangGraph', level=1)
    doc.add_paragraph('LangGraph là một thư viện giúp xây dựng các ứng dụng LLM phức tạp, có trạng thái.')
    doc.add_paragraph('Nó cho phép tạo các đồ thị có chu trình, rất phù hợp cho các tác vụ lặp đi lặp lại như agent.')
    
    # Section 2: Also short, should be one slide
    doc.add_heading('Slide 2: Các lợi ích chính', level=1)
    doc.add_paragraph('- Dễ dàng quản lý trạng thái qua các bước.')
    doc.add_paragraph('- Hỗ trợ các chu trình và vòng lặp một cách tự nhiên.')
    doc.add_paragraph('- Tăng khả năng gỡ lỗi và quan sát luồng dữ liệu.')
    
    # Section 3: Long, should be split into TWO slides
    doc.add_heading('Slide 3: Ví dụ về nội dung dài', level=1)
    long_text_part1 = "Phần đầu tiên của nội dung này đủ ngắn để vừa vặn. Nó giới thiệu khái niệm và thiết lập bối cảnh. Chúng ta sẽ thảo luận về các nguyên tắc cơ bản và lý do tại sao phương pháp này lại hiệu quả. Đoạn này bao gồm một vài câu để đảm bảo nó có nội dung đáng kể, nhưng không quá dài đến mức tự nó vượt quá giới hạn."
    long_text_part2 = "Bây giờ, phần thứ hai này là phần đẩy nội dung vượt quá giới hạn 600 ký tự. Bằng cách thêm đoạn văn này, tổng độ dài của 'Slide 3' sẽ yêu cầu logic phân tách phải được kích hoạt. Trình xử lý thông minh sẽ giữ 'Phần đầu tiên' trên một slide và đặt 'Phần thứ hai' này trên một slide mới, nhưng cả hai đều có cùng một tiêu đề 'Slide 3'."
    doc.add_paragraph(long_text_part1)
    doc.add_paragraph(long_text_part2)

    doc.save(sample_docx_path)
    print(f"Created test file at: {sample_docx_path}\n")

    initial_state = {'docx_file': sample_docx_path}
    result_state = initialize_and_prepare(initial_state)

    # --- Print the final output for verification ---
    final_chunks = result_state["slide_chunks"]
    print("\n==================== VERIFICATION ====================")
    print(f"The code produced a total of {len(final_chunks)} final chunks (slides).")
    print("======================================================")

    for i, chunk in enumerate(final_chunks):
        print(f"\n--- CHUNK #{i+1} (Length: {len(chunk)}) ---")
        print(chunk)
        print("--------------------------------------------------")

    # --- Cleanup ---
    print("\nCleaning up generated files...")
    os.remove(sample_docx_path)
    os.rmdir(result_state['temp_dir'])
    print("Done.")