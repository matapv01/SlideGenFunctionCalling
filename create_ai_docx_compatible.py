from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

OUT = Path("triet_tue_nhan_tao_ai_slidegen.docx")


SECTIONS = [
    (
        "1. Tổng quan về Trí tuệ nhân tạo",
        "Trí tuệ nhân tạo (Artificial Intelligence – AI) là lĩnh vực của khoa học máy tính tập trung vào việc xây dựng các hệ thống có khả năng thực hiện những nhiệm vụ thường cần đến sự thông minh của con người. Các khả năng phổ biến gồm học từ dữ liệu, nhận diện hình ảnh, hiểu ngôn ngữ, suy luận, dự đoán và ra quyết định.",
        "AI ngày càng phát triển nhờ sự gia tăng của dữ liệu, năng lực tính toán và các thuật toán học máy hiện đại. Từ các hệ thống đề xuất nội dung đến trợ lý ảo và xe tự hành, AI đang trở thành một thành phần quan trọng trong nhiều lĩnh vực.",
    ),
    (
        "2. Các công nghệ cốt lõi của AI",
        "Một số công nghệ nền tảng của AI gồm Machine Learning (học máy), Deep Learning (học sâu), Computer Vision (thị giác máy tính) và Natural Language Processing (xử lý ngôn ngữ tự nhiên).",
        "Machine Learning cho phép mô hình học quy luật từ dữ liệu thay vì phải lập trình mọi quy tắc thủ công. Deep Learning sử dụng các mạng nơ-ron nhiều lớp để xử lý dữ liệu phức tạp. Computer Vision tập trung vào hình ảnh và video, trong khi NLP giúp máy tính phân tích và sinh ngôn ngữ tự nhiên.",
    ),
    (
        "3. Quy trình xây dựng một hệ thống AI",
        "Một hệ thống AI thường bắt đầu từ việc xác định bài toán và thu thập dữ liệu. Dữ liệu sau đó được làm sạch, gán nhãn hoặc biến đổi để phù hợp với mô hình. Tiếp theo, dữ liệu được chia thành các tập huấn luyện, kiểm thử và đánh giá.",
        "Sau khi lựa chọn thuật toán, mô hình được huấn luyện và tối ưu các tham số. Cuối cùng, mô hình được đánh giá bằng các chỉ số phù hợp và triển khai vào ứng dụng thực tế. Trong quá trình vận hành, hệ thống cần được theo dõi và cập nhật khi dữ liệu hoặc yêu cầu thay đổi.",
    ),
    (
        "4. Ứng dụng của AI trong thực tế",
        "AI được ứng dụng rộng rãi trong y tế, tài chính, giáo dục, sản xuất, giao thông và thương mại. Trong y tế, AI có thể hỗ trợ phân tích ảnh y khoa và dự đoán nguy cơ bệnh. Trong tài chính, AI được dùng để phát hiện giao dịch bất thường và đánh giá rủi ro.",
        "Trong doanh nghiệp, AI còn hỗ trợ chatbot, tìm kiếm thông tin, phân tích tài liệu, tự động hóa quy trình và tạo nội dung. Generative AI cho phép tạo văn bản, hình ảnh, âm thanh và mã nguồn dựa trên yêu cầu của người dùng.",
    ),
    (
        "5. Thách thức và xu hướng phát triển",
        "Bên cạnh tiềm năng lớn, AI vẫn tồn tại các vấn đề như chất lượng dữ liệu, độ chính xác, chi phí tính toán, bảo mật và quyền riêng tư. Mô hình cũng có thể tạo ra kết quả sai hoặc mang thiên kiến nếu dữ liệu và quy trình đánh giá không phù hợp.",
        "Trong tương lai, AI được dự đoán tiếp tục phát triển theo hướng mô hình đa phương thức, AI Agent, hệ thống RAG, mô hình nhỏ chạy trên thiết bị biên và các giải pháp AI có khả năng phối hợp nhiều công cụ. Việc phát triển AI có trách nhiệm, minh bạch và an toàn sẽ ngày càng quan trọng.",
    ),
]


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Arial"
    heading1.font.size = Pt(14)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(15, 118, 110)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Arial"
    heading2.font.size = Pt(12)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(23, 60, 63)


def add_title(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("TRÍ TUỆ NHÂN TẠO (AI)")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(15, 118, 110)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("Tổng quan, công nghệ và ứng dụng")
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(82, 101, 104)


def add_chart_data_table(document: Document) -> None:
    document.add_heading("Bảng dữ liệu đầu vào cho SlideGen", level=2)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    header[0].text = "Category (Năm)"
    header[1].text = "Machine Learning"
    header[2].text = "Generative AI"
    header[3].text = "AI Agent"
    values = [
        (2020, 18, 8, 3),
        (2021, 29, 13, 5),
        (2022, 41, 24, 9),
        (2023, 55, 39, 18),
        (2024, 69, 58, 31),
        (2025, 82, 76, 49),
    ]
    for year, machine_learning, generative_ai, ai_agent in values:
        cells = table.add_row().cells
        cells[0].text = str(year)
        cells[1].text = str(machine_learning)
        cells[2].text = str(generative_ai)
        cells[3].text = str(ai_agent)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
                    if row is table.rows[0]:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    run = note.add_run("Ghi chú: Số liệu trong bảng là dữ liệu minh họa, không đại diện cho một nguồn thống kê chính thức.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(82, 101, 104)


def build_document() -> None:
    document = Document()
    configure_styles(document)
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    document.core_properties.title = "Trí tuệ nhân tạo (AI) – Tổng quan, công nghệ và ứng dụng"
    document.core_properties.subject = "Tổng quan, công nghệ, ứng dụng và xu hướng AI"
    document.core_properties.author = "OpenAI"

    add_title(document)
    for title, first, second in SECTIONS:
        document.add_heading(title, level=1)
        document.add_paragraph(first)
        document.add_paragraph(second)

    document.add_page_break()
    document.add_heading("6. Xu hướng gia tăng mức độ ứng dụng AI", level=1)
    document.add_paragraph(
        "Cùng với sự trưởng thành của hạ tầng dữ liệu, điện toán đám mây và các mô hình nền tảng, mức độ ứng dụng AI có xu hướng tăng qua thời gian. Doanh nghiệp thường bắt đầu từ các bài toán có phạm vi rõ ràng như chatbot, tự động hóa quy trình và phân tích dữ liệu, sau đó mở rộng sang các hệ thống hỗ trợ ra quyết định và Generative AI."
    )
    document.add_paragraph(
        "Đây là dữ liệu phù hợp để SlideGen tạo biểu đồ cột nhóm theo thời gian. Khi sử dụng trong báo cáo chính thức, có thể thay bằng dữ liệu khảo sát hoặc số liệu doanh nghiệp thực tế."
    )
    tool_hint = document.add_paragraph()
    tool_hint.paragraph_format.space_before = Pt(4)
    tool_hint.paragraph_format.space_after = Pt(8)
    hint = tool_hint.add_run(
        "Gợi ý cho SlideGen — tool: generate_best_selling_product_chart_slide_xml | "
        "category: Năm | series: Machine Learning, Generative AI, AI Agent | "
        "chart type: grouped column chart"
    )
    hint.bold = True
    hint.font.size = Pt(10)
    hint.font.color.rgb = RGBColor(15, 118, 110)
    add_chart_data_table(document)

    document.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build_document()
