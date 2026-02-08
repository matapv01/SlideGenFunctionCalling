"""
SlideGen Backend - Main FastAPI Application
"""
from fastapi import FastAPI, File, UploadFile, HTTPException, Request, Form, Depends, Query
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
import logging
import os
import shutil
import zipfile
import json
import re
import asyncio
from pathlib import Path
from datetime import datetime
from typing import AsyncGenerator

from .config import (
    PROJECT_DIR, STATIC_DIR, FRONTEND_DIR, IMAGE_DIR, UPLOAD_DIR, 
    OUTPUT_DIR, TEMP_DIR, CAPTURES_DIR, CORS_ORIGINS, 
    DISABLE_CONVERTER, LOG_LEVEL
)
from .schemas import (
    CaptureSlideRequest, CaptureMultipleSlidesRequest,
    ProjectUploadResponse, ConversionStatusResponse,
    HealthCheckResponse, StructureUpdateRequest, StructureResponse,
    StructureApproveRequest
)
from .services import SlideService, CaptureService, ConversionService, WorkflowService

# Import Pixabay tool and Translation tool
import sys
from pathlib import Path as SysPath
project_root = SysPath(__file__).parent.parent.parent.parent
sys.path.insert(0, str(project_root))
from src.tools.pixabay_tool import PixabayTool
from src.tools.translate_tool import get_translator

# Setup logging
logging.basicConfig(level=getattr(logging, LOG_LEVEL))
logger = logging.getLogger(__name__)

# Initialize FastAPI app
app = FastAPI(
    title="SlideGen API",
    description="API for slide generation and management",
    version="1.0.0"
)

# CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["*"]
)

# Add middleware for Private Network Access (Chrome security policy)
@app.middleware("http")
async def add_private_network_header(request: Request, call_next):
    response = await call_next(request)
    # Allow requests from public networks to private networks (localhost)
    response.headers["Access-Control-Allow-Private-Network"] = "true"
    return response

# Middleware to handle clean URLs and block /frontend/*.html direct access
@app.middleware("http")
async def clean_url_middleware(request: Request, call_next):
    path = request.url.path
    query_string = request.url.query
    
    # Map for blocking /frontend/*.html and redirecting to clean URLs
    frontend_html_redirects = {
        "/frontend/index.html": "/",
        "/frontend/login.html": "/login",
        "/frontend/register.html": "/register",
        "/frontend/projects.html": "/projects",
        "/frontend/edit.html": "/edit",
        "/frontend/guide.html": "/guide",
        "/frontend/streaming.html": "/streaming",
        "/frontend/structure-editor.html": "/structure-editor",
        "/frontend/view_project.html": "/view-project",
    }
    
    # Redirect /frontend/*.html to clean URL (preserve query string)
    if path in frontend_html_redirects:
        from fastapi.responses import RedirectResponse
        new_url = frontend_html_redirects[path]
        if query_string:
            new_url = f"{new_url}?{query_string}"
        return RedirectResponse(url=new_url, status_code=301)
    
    # Handle root-level *.html requests (from relative links in HTML)
    # e.g., /login.html -> /login, /projects.html -> /projects
    root_html_files = {
        "/index.html": "/",
        "/login.html": "/login",
        "/register.html": "/register",
        "/projects.html": "/projects",
        "/edit.html": "/edit",
        "/guide.html": "/guide",
        "/streaming.html": "/streaming",
        "/structure-editor.html": "/structure-editor",
        "/view_project.html": "/view-project",
    }
    
    # Redirect /*.html to clean URL (preserve query string)
    if path in root_html_files:
        from fastapi.responses import RedirectResponse
        new_url = root_html_files[path]
        if query_string:
            new_url = f"{new_url}?{query_string}"
        return RedirectResponse(url=new_url, status_code=302)
    
    # Serve frontend assets from root level (config.js, css/, js/)
    # This allows relative paths like ./config.js to work from clean URLs
    if path in ["/config.js", "/script.js"]:
        file_path = FRONTEND_DIR / path.lstrip("/")
        if file_path.exists():
            return FileResponse(str(file_path), media_type="application/javascript")
    
    # Serve css/ and js/ directories from root
    if path.startswith("/css/") or path.startswith("/js/"):
        file_path = FRONTEND_DIR / path.lstrip("/")
        if file_path.exists():
            # Determine content type based on extension
            if path.endswith(".css"):
                media_type = "text/css"
            elif path.endswith(".js"):
                media_type = "application/javascript"
            else:
                media_type = None
            return FileResponse(str(file_path), media_type=media_type)
    
    return await call_next(request)

# Initialize database
from .database import init_db
init_db()

# Mount static files
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
# Mount frontend for CSS, JS and other assets (HTML blocked by middleware)
app.mount("/frontend", StaticFiles(directory=str(FRONTEND_DIR)), name="frontend")

# Initialize services
slide_service = SlideService(STATIC_DIR, IMAGE_DIR)

# Initialize conversion service
conversion_service = None
if not DISABLE_CONVERTER:
    try:
        conversion_service = ConversionService(PROJECT_DIR)
        if conversion_service.is_available():
            logger.info('✅ ConversionService initialized successfully')
        else:
            logger.warning('⚠️  ConversionService dependencies not available')
            conversion_service = None
    except Exception as e:
        logger.warning(f'Failed to initialize ConversionService: {e}')
        conversion_service = None
else:
    logger.info('DISABLE_CONVERTER set: converter functionality is disabled')

# Initialize capture service
capture_service = None
try:
    # Import capture module dynamically
    import sys
    sys.path.insert(0, str(PROJECT_DIR))
    from capture import SlideCapture
    
    slide_capture = SlideCapture()
    capture_service = CaptureService(CAPTURES_DIR, slide_capture)
    logger.info('✅ CaptureService initialized successfully')
except Exception as e:
    logger.warning(f'Failed to initialize CaptureService: {e}')
    capture_service = None

# Initialize workflow service
workflow_service = None
try:
    workflow_service = WorkflowService(str(PROJECT_DIR))
    if workflow_service.is_available():
        logger.info('✅ WorkflowService initialized successfully')
    else:
        logger.warning('⚠️  WorkflowService not available - using fallback slides')
except Exception as e:
    logger.warning(f'Failed to initialize WorkflowService: {e}')
    workflow_service = None


# In-memory storage for structure editor sessions
# TODO: Replace with Redis for production
structure_sessions = {}


# ========================== UTILITY FUNCTIONS ==========================

def extract_slide_index(path):
    """
    Extract slide index from filename like:
    slide_10_0.html → 10
    """
    match = re.search(r"slide_(\d+)", path.stem)
    return int(match.group(1)) if match else float("inf")

def save_workflow_to_database(
    user_id: int,
    project_name: str,
    source_type: str,
    source_data: dict,
    slide_options: dict,
    output_folder: Path
) -> int:
    """Save completed workflow to database with converted HTML"""
    try:
        # Create project
        project_id = create_project(
            user_id=user_id,
            title=project_name,
            description=source_data.get('description', ''),
            source_type=source_type,
            source_data=source_data
        )
        
        # Save project metadata
        save_project_metadata(
            project_id=project_id,
            slide_options=slide_options,
            workflow_config={},
            generation_params={}
        )
        
        # Convert HTML slides to absolute positioning before saving
        html_folder = output_folder / "html"
        converted_folder = output_folder / "converted"
        
        if html_folder.exists():
            # Convert HTML slides using converter.js
            logger.info(f"🔄 Converting HTML slides to absolute positioning...")
            conversion_result = conversion_service.convert_slides_batch(html_folder, converted_folder)
            
            if conversion_result['success']:
                logger.info(f"✅ Successfully converted {len(conversion_result['success'])} slides")
            if conversion_result['failed']:
                logger.warning(f"⚠️ Failed to convert {len(conversion_result['failed'])} slides")
            
            # Read and save converted slides (prioritize converted, fallback to original)
            html_files = sorted(
                html_folder.glob("*.html"),
                key=extract_slide_index
            )
            
            for idx, html_file in enumerate(html_files):
                # Try to read converted HTML first
                converted_file = converted_folder / f"converted_{html_file.name}"
                if converted_file.exists():
                    logger.info(f"📄 Reading converted HTML: {converted_file.name}")
                    with open(converted_file, 'r', encoding='utf-8') as f:
                        html_content = f.read()
                else:
                    # Fallback to original HTML if conversion failed
                    logger.warning(f"⚠️ Using original HTML (conversion failed): {html_file.name}")
                    with open(html_file, 'r', encoding='utf-8') as f:
                        html_content = f.read()
                
                # Extract slide type from filename
                slide_name = html_file.stem
                slide_type = 'content'  # Default
                
                if 'intro' in slide_name.lower() or idx == 0:
                    slide_type = 'intro'
                elif 'transition' in slide_name.lower():
                    slide_type = 'transition'
                elif 'thank' in slide_name.lower():
                    slide_type = 'thank_you'
                
                # Create slide in database with converted HTML
                create_slide(
                    project_id=project_id,
                    slide_index=idx,
                    slide_type=slide_type,
                    title=f"Slide {idx + 1}",
                    html_content=html_content,
                    design_settings={}
                )
            
            # Update project slide count and status
            update_project(
                project_id=project_id,
                slide_count=len(html_files),
                status='completed'
            )
            
            logger.info(f"✅ Saved project {project_id} with {len(html_files)} slides to database")
        
        return project_id
        
    except Exception as e:
        logger.error(f"Failed to save workflow to database: {e}")
        return 0

def clear_temp_dir(temp_dir: Path):
    """Clear temporary directory"""
    if temp_dir.exists():
        for item in temp_dir.iterdir():
            try:
                if item.is_file() or item.is_symlink():
                    item.unlink()
                elif item.is_dir():
                    shutil.rmtree(item)
            except Exception as e:
                logger.warning(f"Could not remove {item}: {e}")


# ========================== ROOT ENDPOINTS ==========================

@app.get("/")
async def read_root():
    """Serve frontend index page"""
    return FileResponse(str(FRONTEND_DIR / "index.html"), media_type="text/html")


# Short URL routes for cleaner access
@app.get("/login")
async def login_page():
    """Serve login page"""
    return FileResponse(str(FRONTEND_DIR / "login.html"), media_type="text/html")


@app.get("/register")
async def register_page():
    """Serve register page"""
    return FileResponse(str(FRONTEND_DIR / "register.html"), media_type="text/html")


@app.get("/projects")
async def projects_page():
    """Serve projects page"""
    return FileResponse(str(FRONTEND_DIR / "projects.html"), media_type="text/html")


@app.get("/edit")
async def edit_page():
    """Serve edit page"""
    return FileResponse(str(FRONTEND_DIR / "edit.html"), media_type="text/html")


@app.get("/guide")
async def guide_page():
    """Serve guide page"""
    return FileResponse(str(FRONTEND_DIR / "guide.html"), media_type="text/html")


@app.get("/streaming")
async def streaming_page():
    """Serve streaming page"""
    return FileResponse(str(FRONTEND_DIR / "streaming.html"), media_type="text/html")


@app.get("/structure-editor")
async def structure_editor_page():
    """Serve structure editor page"""
    return FileResponse(str(FRONTEND_DIR / "structure-editor.html"), media_type="text/html")


@app.get("/view-project")
async def view_project_page():
    """Serve view project page"""
    return FileResponse(str(FRONTEND_DIR / "view_project.html"), media_type="text/html")


@app.get("/api")
async def api_info():
    """API Root - Returns API information"""
    return {
        "name": "SlideGen Backend API",
        "version": "1.0.0",
        "status": "running",
        "docs": "/docs"
    }


# ========================== AUTHENTICATION ENDPOINTS ==========================

from .models import UserCreate, UserLogin, Token, UserResponse, User
from .auth import (
    get_password_hash, 
    verify_password,
    authenticate_user, 
    create_access_token,
    get_current_active_user
)
from .database import get_user_by_email, create_user, get_db
from datetime import timedelta

@app.post("/api/auth/register", response_model=UserResponse, status_code=201)
async def register(user_data: UserCreate):
    """Register a new user"""
    # Check if user already exists
    existing_user = get_user_by_email(user_data.email)
    if existing_user:
        raise HTTPException(
            status_code=400,
            detail="Email already registered"
        )
    
    # Hash password and create user
    hashed_password = get_password_hash(user_data.password)
    user_id = create_user(user_data.email, user_data.username, hashed_password)
    
    # Get created user
    user = get_user_by_email(user_data.email)
    
    return UserResponse(
        id=user["id"],
        email=user["email"],
        username=user["username"],
        created_at=user["created_at"],
        is_active=bool(user["is_active"]),
        is_vip=bool(user.get("is_vip", False))
    )

@app.post("/api/auth/login", response_model=Token)
async def login(user_data: UserLogin):
    """Login user and return JWT token"""
    user = authenticate_user(user_data.email, user_data.password)
    if not user:
        raise HTTPException(
            status_code=401,
            detail="Incorrect email or password"
        )
    
    # Create access token
    access_token_expires = timedelta(minutes=60 * 24)  # 24 hours
    access_token = create_access_token(
        data={"sub": user["email"]},
        expires_delta=access_token_expires
    )
    
    return Token(access_token=access_token, token_type="bearer")

@app.get("/api/auth/me", response_model=UserResponse)
async def get_me(current_user: UserResponse = Depends(get_current_active_user)):
    """Get current user information"""
    return current_user

@app.post("/api/auth/logout")
async def logout():
    """Logout user (client-side token removal)"""
    return {"message": "Successfully logged out"}


@app.put("/api/auth/change-password")
async def change_password(
    password_data: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Change user password"""
    current_password = password_data.get('current_password')
    new_password = password_data.get('new_password')
    
    if not current_password or not new_password:
        raise HTTPException(status_code=400, detail="Missing required fields")
    
    # Get full user data from database (including hashed_password)
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM users WHERE id = ?', (current_user.id,))
        user_row = cursor.fetchone()
        
        if not user_row:
            raise HTTPException(status_code=404, detail="User not found")
        
        # Verify current password
        if not verify_password(current_password, user_row['hashed_password']):
            raise HTTPException(status_code=400, detail="Current password is incorrect")
        
        # Validate new password
        if len(new_password) < 6:
            raise HTTPException(status_code=400, detail="New password must be at least 6 characters")
        
        # Update password in database
        try:
            new_hashed_password = get_password_hash(new_password)
            cursor.execute(
                'UPDATE users SET hashed_password = ? WHERE id = ?',
                (new_hashed_password, current_user.id)
            )
            conn.commit()
            return {"message": "Password changed successfully"}
        except Exception as e:
            conn.rollback()
            raise HTTPException(status_code=500, detail=f"Failed to change password: {str(e)}")


# ========================== USER VIP ENDPOINTS ==========================

from .database import update_user_vip_status

@app.get("/api/user/me", response_model=UserResponse)
async def get_current_user_info(current_user: UserResponse = Depends(get_current_active_user)):
    """Get current user information including VIP status"""
    return current_user

@app.post("/api/user/upgrade-vip")
async def upgrade_to_vip(current_user: UserResponse = Depends(get_current_active_user)):
    """Upgrade current user to VIP status"""
    try:
        update_user_vip_status(current_user.id, True)
        logger.info(f"✅ User {current_user.email} upgraded to VIP")
        return {
            "message": "Successfully upgraded to VIP!",
            "is_vip": True
        }
    except Exception as e:
        logger.error(f"Failed to upgrade user to VIP: {e}")
        raise HTTPException(status_code=500, detail="Failed to upgrade to VIP")



# ========================== USER PROFILE & SETTINGS ==========================

from .database import (
    get_user_profile, create_or_update_user_profile,
    get_user_settings, create_default_settings, update_user_settings,
    create_project, get_user_projects, get_project, update_project, delete_project,
    create_slide, get_project_slides as db_get_project_slides, get_slide, 
    update_slide, delete_slide, save_project_metadata, get_project_metadata
)

@app.get("/api/profile")
async def get_profile(current_user: UserResponse = Depends(get_current_active_user)):
    """Get current user profile"""
    profile = get_user_profile(current_user.id)
    if not profile:
        # Return empty profile
        return {
            "user_id": current_user.id,
            "full_name": None,
            "avatar_url": None,
            "bio": None
        }
    return profile

@app.put("/api/profile")
async def update_profile(
    profile_data: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Update user profile"""
    create_or_update_user_profile(
        user_id=current_user.id,
        full_name=profile_data.get("full_name"),
        avatar_url=profile_data.get("avatar_url"),
        bio=profile_data.get("bio")
    )
    return {"message": "Profile updated"}

@app.get("/api/settings")
async def get_settings(current_user: UserResponse = Depends(get_current_active_user)):
    """Get user settings"""
    settings = get_user_settings(current_user.id)
    if not settings:
        # Create default settings
        create_default_settings(current_user.id)
        settings = get_user_settings(current_user.id)
    return settings

@app.put("/api/settings")
async def update_settings(
    settings_data: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Update user settings"""
    update_user_settings(
        user_id=current_user.id,
        theme=settings_data.get("theme"),
        default_transition_enabled=settings_data.get("default_transition_enabled"),
        default_thank_you_enabled=settings_data.get("default_thank_you_enabled"),
        default_content_length=settings_data.get("default_content_length"),
        language=settings_data.get("language"),
        settings_json=settings_data.get("settings_json")
    )
    return {"message": "Settings updated"}


# ========================== PROJECTS API ==========================

@app.get("/api/projects")
async def list_projects(
    limit: int = 50,
    offset: int = 0,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Get all projects for current user"""
    projects = get_user_projects(current_user.id, limit, offset)
    return {
        "projects": projects,
        "total": len(projects)
    }

@app.post("/api/projects")
async def create_new_project(
    project_data: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Create new project"""
    project_id = create_project(
        user_id=current_user.id,
        title=project_data.get("title", "Untitled Project"),
        description=project_data.get("description"),
        source_type=project_data.get("source_type", "idea"),
        source_data=project_data.get("source_data")
    )
    return {"project_id": project_id, "message": "Project created"}

@app.get("/api/projects/{project_id}")
async def get_project_detail(
    project_id: int,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Get project details with slides"""
    project = get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Verify ownership
    if project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    # Get slides
    slides = db_get_project_slides(project_id)
    
    # Get metadata
    metadata = get_project_metadata(project_id)
    
    return {
        "project": project,
        "slides": slides,
        "metadata": metadata
    }

@app.put("/api/projects/{project_id}")
async def update_project_detail(
    project_id: int,
    project_data: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Update project"""
    project = get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Verify ownership
    if project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    update_project(
        project_id=project_id,
        title=project_data.get("title"),
        description=project_data.get("description"),
        status=project_data.get("status"),
        slide_count=project_data.get("slide_count"),
        thumbnail_url=project_data.get("thumbnail_url")
    )
    return {"message": "Project updated"}

@app.delete("/api/projects/{project_id}")
async def delete_project_detail(
    project_id: int,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Delete project"""
    project = get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Verify ownership
    if project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    delete_project(project_id)
    return {"message": "Project deleted"}


# ========================== SLIDES API (DATABASE) ==========================

@app.get("/api/projects/{project_id}/slides")
async def get_project_slides_list(
    project_id: int,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Get all slides for a project"""
    project = get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Verify ownership
    if project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    slides = db_get_project_slides(project_id)
    return {"slides": slides, "total": len(slides)}

@app.post("/api/projects/{project_id}/slides")
async def create_new_slide(
    project_id: int,
    slide_data: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Create a new slide in a project"""
    project = get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    
    # Verify ownership
    if project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    # Create slide
    slide_id = create_slide(
        project_id=project_id,
        slide_index=slide_data.get("slide_index", 0),
        slide_type=slide_data.get("slide_type", "content"),
        title=slide_data.get("title"),
        content_text=slide_data.get("content_text"),
        html_content=slide_data.get("html_content"),
        design_settings=slide_data.get("design_settings")
    )
    
    # Update project slide count
    slides = db_get_project_slides(project_id)
    update_project(project_id, slide_count=len(slides))
    
    return {"id": slide_id, "message": "Slide created"}

@app.get("/api/slides/{slide_id}")
async def get_slide_detail(
    slide_id: int,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Get slide details"""
    slide = get_slide(slide_id)
    if not slide:
        raise HTTPException(status_code=404, detail="Slide not found")
    
    # Verify ownership through project
    project = get_project(slide["project_id"])
    if not project or project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    return slide

@app.put("/api/slides/{slide_id}")
async def update_slide_content(
    slide_id: int,
    slide_data: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Update slide content"""
    slide = get_slide(slide_id)
    if not slide:
        raise HTTPException(status_code=404, detail="Slide not found")
    
    # Verify ownership through project
    project = get_project(slide["project_id"])
    if not project or project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    update_slide(
        slide_id=slide_id,
        title=slide_data.get("title"),
        content_text=slide_data.get("content_text"),
        html_content=slide_data.get("html_content"),
        design_settings=slide_data.get("design_settings")
    )
    return {"message": "Slide updated"}

@app.delete("/api/slides/{slide_id}")
async def delete_slide_detail(
    slide_id: int,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Delete a slide"""
    slide = get_slide(slide_id)
    if not slide:
        raise HTTPException(status_code=404, detail="Slide not found")
    
    # Verify ownership through project
    project = get_project(slide["project_id"])
    if not project or project["user_id"] != current_user.id:
        raise HTTPException(status_code=403, detail="Access denied")
    
    delete_slide(slide_id)
    
    # Update project slide count
    slides = db_get_project_slides(slide["project_id"])
    update_project(slide["project_id"], slide_count=len(slides))
    
    return {"message": "Slide deleted"}


# ========================== UPLOAD ENDPOINTS ==========================

@app.post("/api/upload-docx")
async def upload_docx(
    file: UploadFile = File(...),
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Upload and process DOCX file with AI workflow (Protected route)"""
    if not file.filename.endswith('.docx'):
        raise HTTPException(status_code=400, detail="Only DOCX files accepted")
    
    # Create project name from filename
    project_name = os.path.splitext(file.filename)[0]
    output_folder = OUTPUT_DIR / project_name
    
    # Handle existing project folder
    if output_folder.exists():
        try:
            shutil.rmtree(output_folder)
        except PermissionError:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            project_name = f"{project_name}_{timestamp}"
            output_folder = OUTPUT_DIR / project_name
    
    output_folder.mkdir(parents=True, exist_ok=True)
    
    # Save temp file
    temp_path = TEMP_DIR / file.filename
    with open(temp_path, "wb") as f:
        content = await file.read()
        f.write(content)
    
    try:
        # Try to process with AI workflow first
        if workflow_service and workflow_service.is_available():
            logger.info(f"🤖 Processing DOCX with AI workflow: {file.filename}")
            
            result_zip = await workflow_service.process_docx_with_workflow(
                docx_path=str(temp_path),
                output_folder=str(output_folder),
                router_flags=False
            )
            
            if result_zip and os.path.exists(result_zip):
                logger.info(f"✅ Successfully processed with workflow")
                
                # Extract ZIP to output folder
                html_folder = output_folder / "html"
                html_folder.mkdir(exist_ok=True)
                
                try:
                    with zipfile.ZipFile(result_zip, 'r') as zip_ref:
                        for file_info in zip_ref.infolist():
                            # Skip directories and temp_debug
                            if file_info.filename.endswith('/') or 'temp_debug' in file_info.filename:
                                continue
                            try:
                                zip_ref.extract(file_info, output_folder)
                            except Exception as e:
                                logger.warning(f"Error extracting {file_info.filename}: {e}")
                    
                    logger.info(f"Successfully extracted workflow results")
                except Exception as e:
                    logger.error(f"Error extracting ZIP: {e}")
                    raise
            else:
                raise Exception("Workflow did not produce valid output")
        else:
            logger.warning("⚠️ Workflow not available, using sample slides")
            raise Exception("Workflow not available")
            
    except Exception as e:
        # Fallback: use sample HTML slides
        logger.warning(f"⚠️ Fallback to sample slides: {e}")
        
        html_folder = output_folder / "html"
        html_folder.mkdir(exist_ok=True)
        
        # Copy sample HTML files
        sample_html_dir = STATIC_DIR / "sample_html"
        if sample_html_dir.exists():
            for html_file in sample_html_dir.glob("*.html"):
                shutil.copy2(html_file, html_folder / html_file.name)
        
    finally:
        if temp_path.exists():
            temp_path.unlink()
    
    # Convert slides if converter available
    if conversion_service and conversion_service.is_available():
        html_folder = output_folder / "html"
        converted_folder = output_folder / "converted"
        
        if html_folder.exists():
            try:
                converted_folder.mkdir(exist_ok=True)
                conversion_result = conversion_service.convert_slides_batch(html_folder, converted_folder)
                logger.info(f"Conversion: {len(conversion_result['success'])} success, {len(conversion_result['failed'])} failed")
            except Exception as e:
                logger.error(f"Conversion error: {e}")
                logger.error(f"Conversion error: {e}")
    
    # Build slides response
    slides = slide_service.build_slides_from_output(str(output_folder))
    
    return JSONResponse(content={"slides": slides, "project_name": project_name})


@app.post("/api/upload-docx-for-stream")
async def upload_docx_for_stream(
    file: UploadFile = File(...),
    enable_transition_slides: str = Form("true"),
    enable_thank_you_slide: str = Form("true"),
    workflow_mode: str = Form("auto")
):
    """Upload DOCX file and return filename for streaming"""
    if not file.filename.endswith('.docx'):
        raise HTTPException(status_code=400, detail="Only DOCX files accepted")
    
    # Clean up old temp files first
    try:
        deleted_count = 0
        for old_file in TEMP_DIR.glob("*"):
            if old_file.is_file():
                try:
                    old_file.unlink()
                    deleted_count += 1
                except Exception as e:
                    logger.warning(f"Could not delete {old_file}: {e}")
        if deleted_count > 0:
            logger.info(f"🗑️ Cleaned up {deleted_count} old temp files")
    except Exception as e:
        logger.warning(f"Error cleaning temp directory: {e}")
    
    # Convert string to boolean
    def str_to_bool(value: str) -> bool:
        if isinstance(value, bool):
            return value
        return value.lower() in ('true', '1', 'yes', 'on')
    
    logger.info(f"📥 Received form data:")
    logger.info(f"   - enable_transition_slides (raw): '{enable_transition_slides}' (type: {type(enable_transition_slides)})")
    logger.info(f"   - enable_thank_you_slide (raw): '{enable_thank_you_slide}' (type: {type(enable_thank_you_slide)})")
    
    enable_transition = str_to_bool(enable_transition_slides)
    enable_thank_you = str_to_bool(enable_thank_you_slide)
    
    logger.info(f"🔄 After parsing:")
    logger.info(f"   - enable_transition: {enable_transition}")
    logger.info(f"   - enable_thank_you: {enable_thank_you}")
    
    # Save temp file
    temp_path = TEMP_DIR / file.filename
    with open(temp_path, "wb") as f:
        content = await file.read()
        f.write(content)
    
    # Save options to JSON file with same name
    import json
    options_path = TEMP_DIR / f"{file.filename}.options.json"
    
    # Delete old options file if exists to avoid caching
    if options_path.exists():
        try:
            options_path.unlink()
            logger.info(f"🗑️ Deleted old options file: {options_path}")
        except Exception as e:
            logger.warning(f"⚠️ Could not delete old options file: {e}")
    
    # Write new options
    options_data = {
        "enable_transition_slides": enable_transition,
        "enable_thank_you_slide": enable_thank_you,
        "workflow_mode": workflow_mode
    }
    
    try:
        with open(options_path, "w") as f:
            json.dump(options_data, f, indent=2)
        logger.info(f"💾 Successfully saved slide options to {options_path}")
        logger.info(f"📄 File content: {options_data}")
        
        # Verify file was written correctly
        with open(options_path, "r") as f:
            verified = json.load(f)
        logger.info(f"✅ Verified saved options: {verified}")
    except Exception as e:
        logger.error(f"❌ Error saving options: {e}")
    
    return JSONResponse(content={"filename": file.filename})


@app.post("/api/upload-docx-for-manual")
async def upload_docx_for_manual(file: UploadFile = File(...)):
    """Upload DOCX for manual structure editing workflow - Extract plain text for user to manually structure"""
    if not file.filename.endswith('.docx'):
        raise HTTPException(status_code=400, detail="Only DOCX files accepted")
    
    logger.info(f"📝 Manual workflow: Processing {file.filename}")
    
    # Save temp file
    temp_path = TEMP_DIR / file.filename
    with open(temp_path, "wb") as f:
        content = await file.read()
        f.write(content)
    
    try:
        # Extract plain text from DOCX using python-docx
        from docx import Document
        
        doc = Document(temp_path)
        
        # Extract all paragraphs with their positions
        full_text = ""
        paragraphs_data = []
        
        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:  # Skip empty paragraphs
                start_pos = len(full_text)
                full_text += text + "\n\n"
                end_pos = len(full_text)
                
                paragraphs_data.append({
                    "index": idx,
                    "text": text,
                    "start_pos": start_pos,
                    "end_pos": end_pos
                })
        
        logger.info(f"✅ Extracted {len(paragraphs_data)} paragraphs from document")
        logger.info(f"� Total text length: {len(full_text)} characters")
        
        # Create session
        import uuid
        session_id = str(uuid.uuid4())
        
        # Store in structure_sessions with EMPTY structure for user to build
        structure_sessions[session_id] = {
            'document_title': file.filename.replace('.docx', ''),
            'document_text': full_text,  # Full plain text
            'paragraphs': paragraphs_data,  # Paragraph metadata
            'document_structure': [],  # Empty - user will build this
            'docx_file': str(temp_path),
            'filename': file.filename
        }
        
        logger.info(f"✅ Created manual session: {session_id}")
        
        return JSONResponse(content={
            "success": True,
            "session_id": session_id,
            "document_title": file.filename.replace('.docx', ''),
            "text_length": len(full_text),
            "paragraph_count": len(paragraphs_data)
        })
        
    except Exception as e:
        logger.error(f"❌ Error in manual workflow: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to process document: {str(e)}")
    finally:
        # Keep temp file for now (will be deleted after workflow completes)
        pass

from .auth import get_current_user_from_token_param

@app.get("/api/upload-docx-stream")
async def upload_docx_stream(
    filename: str = Query(..., description="DOCX filename to process"),
    token: str = Query(..., description="Authentication token (required)")
):
    """Stream processing of DOCX file with Server-Sent Events - Requires Authentication"""
    
    # Required authentication - validate token
    logger.info(f"🔐 Authenticating request with token: {token[:20] if token else 'None'}...")
    
    # Check if token is missing or is the string "null" (JavaScript encodes null as "null")
    if not token or token == "null" or not token.strip():
        logger.error("❌ Authentication failed: No valid token provided")
        raise HTTPException(
            status_code=401,
            detail="Authentication required. Please log in to generate slides."
        )
    
    # Decode and validate token
    try:
        from .auth import decode_access_token
        token_data = decode_access_token(token)
        from .database import get_user_by_email
        user = get_user_by_email(token_data.email)
        
        if not user or not user.get("is_active"):
            logger.error(f"❌ Authentication failed: User not found or inactive")
            raise HTTPException(
                status_code=401,
                detail="User not found or inactive"
            )
        
        current_user = UserResponse(
            id=user["id"],
            email=user["email"],
            username=user["username"],
            created_at=user["created_at"],
            is_active=bool(user["is_active"]),
            is_vip=bool(user.get("is_vip", False))
        )
        logger.info(f"✅ Authenticated user: {current_user.email}")
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"❌ Authentication failed: {e}")
        raise HTTPException(
            status_code=401,
            detail=f"Invalid or expired token. Please log in again."
        )
    
    # Normalize filename for cross-platform compatibility
    safe_filename = os.path.basename(filename)
    # Preserve Unicode characters (Vietnamese, etc.) but remove dangerous characters
    safe_filename = re.sub(r'[\\/:*?"<>|]', '', safe_filename)
    # Ensure no leading/trailing dots/spaces
    safe_filename = safe_filename.strip('. ')

    temp_path = TEMP_DIR / safe_filename
    if not temp_path.exists() or not safe_filename.lower().endswith('.docx'):
        logger.error(f"DOCX file not found or invalid: {temp_path}")
        raise HTTPException(status_code=400, detail="Invalid file")
    
    # Load slide options
    options_path = TEMP_DIR / f"{safe_filename}.options.json"
    slide_options = {
        "enable_transition_slides": True,
        "enable_thank_you_slide": True
    }
    workflow_mode = "auto"
    
    if options_path.exists():
        try:
            with open(options_path, "r") as f:
                data = json.load(f)
                # Keep compatibility with slide_options format
                slide_options = {
                    "enable_transition_slides": data.get("enable_transition_slides", True),
                    "enable_thank_you_slide": data.get("enable_thank_you_slide", True)
                }
                workflow_mode = data.get("workflow_mode", "auto")
            logger.info(f"📋 Loaded options from file: {slide_options} | mode: {workflow_mode}")
            logger.info(f"   - Transition slides: {slide_options.get('enable_transition_slides')}")
            logger.info(f"   - Thank you slide: {slide_options.get('enable_thank_you_slide')}")
        except Exception as e:
            logger.warning(f"⚠️ Could not load slide options: {e}")
    else:
        logger.warning(f"⚠️ Options file not found: {options_path}")
    
    # Create project name from filename
    project_name = os.path.splitext(safe_filename)[0]
    output_folder = OUTPUT_DIR / project_name
    
    # Handle existing project folder
    if output_folder.exists():
        try:
            shutil.rmtree(output_folder)
        except PermissionError:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            project_name = f"{project_name}_{timestamp}"
            output_folder = OUTPUT_DIR / project_name
    
    output_folder.mkdir(parents=True, exist_ok=True)
    
    # Store temp directory path for cleanup later
    temp_workflow_dir = None
    
    async def event_generator() -> AsyncGenerator[str, None]:
        """Generate Server-Sent Events for workflow progress"""
        nonlocal temp_workflow_dir
        
        # Queue for events
        event_queue = asyncio.Queue()
        workflow_done = asyncio.Event()
        workflow_error = None
        
        def stream_callback(event_type: str, data: dict):
            """Callback to receive events from workflow"""
            try:
                logger.info(f"🔔 Stream callback: {event_type} - {data}")
                
                # For slide_completed, keep the temp path for real-time access
                if event_type == 'slide_completed' and 'html_path' in data:
                    # Keep the original temp path from workflow
                    # Format: /tmp/slidegen_workflow_xxx/output/html/slide_1.html
                    temp_path = data['html_path']
                    # Store as temp path for frontend to access
                    data['html_path'] = temp_path
                    data['is_temp'] = True
                    logger.info(f"🎯 Streaming temp file: {temp_path}")
                
                # Send all events immediately for real-time streaming
                event_queue.put_nowait((event_type, data))
            except Exception as e:
                logger.error(f"Error in stream callback: {e}")
        
        async def run_workflow():
            """Run workflow in background"""
            nonlocal workflow_error, temp_workflow_dir
            try:
                if not workflow_service or not workflow_service.is_available():
                    raise Exception("Workflow service not available")
                
                # Workflow now returns tuple of (result_zip, temp_dir)
                workflow_result = await workflow_service.process_docx_with_workflow(
                    docx_path=str(temp_path),
                    output_folder=str(output_folder),
                    router_flags=False,  # Disable VLM evaluation for all modes
                    callback=stream_callback,
                    slide_options=slide_options,
                    workflow_mode=workflow_mode
                )
                
                if not workflow_result:
                    # Workflow already sent error message via callback
                    logger.error("Workflow returned None (error already sent to frontend)")
                    # Don't return here - let finally block set workflow_done
                else:
                    # Unpack result
                    result_zip, temp_workflow_dir = workflow_result
                    logger.info(f"📁 Temp workflow directory: {temp_workflow_dir}")
                    logger.info(f"📦 Result ZIP: {result_zip}")
                    
                    if result_zip and os.path.exists(result_zip):
                        # Extract ZIP to output folder
                        html_folder = output_folder / "html"
                        html_folder.mkdir(exist_ok=True)
                        
                        with zipfile.ZipFile(result_zip, 'r') as zip_ref:
                            for file_info in zip_ref.infolist():
                                if file_info.filename.endswith('.html'):
                                    zip_ref.extract(file_info, output_folder)
                        
                    if result_zip and os.path.exists(result_zip):
                        # Extract ZIP to output folder
                        html_folder = output_folder / "html"
                        html_folder.mkdir(exist_ok=True)
                        
                        with zipfile.ZipFile(result_zip, 'r') as zip_ref:
                            for file_info in zip_ref.infolist():
                                if not file_info.filename.endswith('/') and 'temp_debug' not in file_info.filename:
                                    try:
                                        zip_ref.extract(file_info, output_folder)
                                    except Exception as e:
                                        logger.warning(f"Error extracting {file_info.filename}: {e}")
                        
                        # Convert slides if converter available
                        if conversion_service and conversion_service.is_available():
                            converted_folder = output_folder / "converted"
                            converted_folder.mkdir(exist_ok=True)
                            conversion_service.convert_slides_batch(html_folder, converted_folder)
                        
                        # Count total slides
                        html_files = list(html_folder.glob("*.html"))
                        total_slides = len(html_files)
                        logger.info(f"📁 Found {total_slides} HTML files")
                        
                        # Save to database
                        project_id = save_workflow_to_database(
                            user_id=current_user.id,
                            project_name=project_name,
                            source_type='docx',
                            source_data={
                                'filename': filename,
                                'description': f'Slides generated from {filename}'
                            },
                            slide_options=slide_options,
                            output_folder=output_folder
                        )
                        
                        # Emit complete event with temp directory path for cleanup
                        await event_queue.put(('complete', {
                            'project_name': project_name,
                            'total_slides': total_slides,
                            'temp_dir': temp_workflow_dir,
                            'project_id': project_id
                        }))
                    else:
                        logger.error("Result ZIP file not found")
                    
            except Exception as e:
                logger.error(f"Workflow error: {e}")
                workflow_error = str(e)
            finally:
                workflow_done.set()
                # Delete uploaded temp file but keep workflow temp directory
                if temp_path.exists():
                    try:
                        temp_path.unlink()
                    except:
                        pass
                # DON'T delete temp_workflow_dir yet - will be cleaned when user goes to editor
        
        # Start workflow in background
        workflow_task = asyncio.create_task(run_workflow())
        
        try:
            # Stream events
            while not workflow_done.is_set() or not event_queue.empty():
                try:
                    # Wait for event with timeout
                    event_type, data = await asyncio.wait_for(event_queue.get(), timeout=0.5)
                    
                    logger.info(f"📤 Sending SSE: {event_type} - {data}")
                    
                    # Format SSE message
                    event_data = json.dumps(data)
                    yield f"event: {event_type}\ndata: {event_data}\n\n"
                    
                except asyncio.TimeoutError:
                    # Send keepalive
                    yield f": keepalive\n\n"
                    continue
            
            # Check for errors
            if workflow_error:
                error_data = json.dumps({"message": workflow_error})
                yield f"event: error_event\ndata: {error_data}\n\n"
            
        except Exception as e:
            logger.error(f"Error in event generator: {e}")
            error_data = json.dumps({"message": str(e)})
            yield f"event: error_event\ndata: {error_data}\n\n"
        finally:
            # Ensure workflow completes
            try:
                await workflow_task
            except:
                pass
    
    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )


@app.get("/api/generate-from-manual-structure-stream")
async def generate_from_manual_structure_stream(
    session_id: str,
    current_user: UserResponse = Depends(get_current_user_from_token_param)
):
    """Stream processing of manual structure from structure editor with SSE"""
    
    # Load approved structure
    approved_file = TEMP_DIR / f"approved_structure_{session_id}.json"
    if not approved_file.exists():
        logger.error(f"Approved structure not found: {approved_file}")
        raise HTTPException(status_code=404, detail="Approved structure not found")
    
    try:
        with open(approved_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except Exception as e:
        logger.error(f"Error loading approved structure: {e}")
        raise HTTPException(status_code=500, detail=f"Error loading structure: {str(e)}")
    
    slide_chunks = data.get('slide_chunks', [])
    document_title = data.get('document_title', 'Untitled')
    
    if not slide_chunks:
        raise HTTPException(status_code=400, detail="No slide chunks found")
    
    logger.info(f"📋 Manual structure session: {session_id}")
    logger.info(f"📄 Document title: {document_title}")
    logger.info(f"📊 Total chunks: {len(slide_chunks)}")
    
    # Create project name
    project_name = document_title.replace(' ', '_')
    output_folder = OUTPUT_DIR / project_name
    
    # Handle existing folder
    if output_folder.exists():
        try:
            shutil.rmtree(output_folder)
        except PermissionError:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            project_name = f"{project_name}_{timestamp}"
            output_folder = OUTPUT_DIR / project_name
    
    output_folder.mkdir(parents=True, exist_ok=True)
    
    temp_workflow_dir = None
    
    async def event_generator() -> AsyncGenerator[str, None]:
        """Generate Server-Sent Events for workflow progress"""
        nonlocal temp_workflow_dir
        
        event_queue = asyncio.Queue()
        workflow_done = asyncio.Event()
        workflow_error = None
        
        def stream_callback(event_type: str, data: dict):
            """Callback to receive events from workflow"""
            try:
                logger.info(f"🔔 Stream callback: {event_type} - {data}")
                
                if event_type == 'slide_completed' and 'html_path' in data:
                    temp_path = data['html_path']
                    data['html_path'] = temp_path
                    data['is_temp'] = True
                    logger.info(f"🎯 Streaming temp file: {temp_path}")
                
                event_queue.put_nowait((event_type, data))
            except Exception as e:
                logger.error(f"Error in stream callback: {e}")
        
        async def run_workflow():
            """Run workflow in background"""
            nonlocal workflow_error, temp_workflow_dir
            try:
                if not workflow_service or not workflow_service.is_available():
                    raise Exception("Workflow service not available")
                
                workflow_result = await workflow_service.process_manual_structure(
                    slide_chunks=slide_chunks,
                    document_title=document_title,
                    output_folder=str(output_folder),
                    callback=stream_callback
                )
                
                if not workflow_result:
                    raise Exception("Workflow did not produce valid output")
                
                result_zip, temp_workflow_dir = workflow_result
                logger.info(f"📁 Temp directory: {temp_workflow_dir}")
                logger.info(f"📦 Result ZIP: {result_zip}")
                
                if result_zip and os.path.exists(result_zip):
                    html_folder = output_folder / "html"
                    html_folder.mkdir(exist_ok=True)
                    
                    with zipfile.ZipFile(result_zip, 'r') as zip_ref:
                        for file_info in zip_ref.infolist():
                            if not file_info.filename.endswith('/') and 'temp_debug' not in file_info.filename:
                                try:
                                    zip_ref.extract(file_info, output_folder)
                                except Exception as e:
                                    logger.warning(f"Error extracting {file_info.filename}: {e}")
                    
                    if conversion_service and conversion_service.is_available():
                        converted_folder = output_folder / "converted"
                        converted_folder.mkdir(exist_ok=True)
                        conversion_service.convert_slides_batch(html_folder, converted_folder)
                    
                    html_files = list(html_folder.glob("*.html"))
                    total_slides = len(html_files)
                    logger.info(f"📁 Found {total_slides} HTML files")
                    
                    project_id = save_workflow_to_database(
                        user_id=current_user.id,
                        project_name=project_name,
                        source_type='manual_structure',
                        source_data={
                            'session_id': session_id,
                            'document_title': document_title,
                            'description': f'Slides from manual structure: {document_title}'
                        },
                        slide_options={'manual_mode': True},
                        output_folder=output_folder
                    )
                    
                    await event_queue.put(('complete', {
                        'project_name': project_name,
                        'total_slides': total_slides,
                        'temp_dir': temp_workflow_dir,
                        'project_id': project_id
                    }))
                else:
                    raise Exception("Workflow did not produce valid output")
                    
            except Exception as e:
                logger.error(f"Workflow error: {e}")
                workflow_error = str(e)
            finally:
                workflow_done.set()
        
        workflow_task = asyncio.create_task(run_workflow())
        
        try:
            while not workflow_done.is_set() or not event_queue.empty():
                try:
                    event_type, data = await asyncio.wait_for(event_queue.get(), timeout=0.5)
                    
                    logger.info(f"📤 Sending SSE: {event_type} - {data}")
                    event_data = json.dumps(data)
                    yield f"event: {event_type}\ndata: {event_data}\n\n"
                    
                except asyncio.TimeoutError:
                    yield f": keepalive\n\n"
                    continue
            
            if workflow_error:
                error_data = json.dumps({"message": workflow_error})
                yield f"event: error_event\ndata: {error_data}\n\n"
            
        except Exception as e:
            logger.error(f"Error in event generator: {e}")
            error_data = json.dumps({"message": str(e)})
            yield f"event: error_event\ndata: {error_data}\n\n"
        finally:
            try:
                await workflow_task
            except:
                pass
    
    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )


# ========================== SLIDES API ==========================

@app.get("/api/projects/by-name/{project_name}/slides")
async def get_project_slides_by_name(project_name: str):
    """Get slides for a project by name - uses converted HTML from database (public access)"""
    logger.info(f"🔓 Public API: Loading slides for project name: {project_name}")
    try:
        # Find project by title/name
        from .database import get_all_projects
        projects = get_all_projects()
        project = next((p for p in projects if p["title"] == project_name), None)
        
        if not project:
            logger.warning(f"Project not found by name: {project_name}")
            raise HTTPException(status_code=404, detail=f"Project '{project_name}' not found")
        
        # Get slides from database (already converted HTML)
        from .database import get_project_slides as get_db_slides
        db_slides = get_db_slides(project["id"])
        
        if not db_slides:
            raise HTTPException(status_code=404, detail=f"No slides found for project: {project_name}")
        
        # Format for frontend
        slides = []
        for slide in db_slides:
            slides.append({
                "id": slide["id"],
                "title": slide["title"],
                "content": slide["html_content"],  # Already converted HTML
                "type": slide["slide_type"],
                "index": slide["slide_index"]
            })
        
        logger.info(f"✅ Loaded {len(slides)} slides from database for {project_name}")
        return JSONResponse(content={
            "project_name": project_name,
            "project_id": project["id"],
            "slides": slides,
            "count": len(slides)
        })
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error getting slides for project {project_name}: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ========================== IMAGE ENDPOINTS ==========================

@app.get("/api/images/{image_name}")
async def get_image(image_name: str):
    """Serve image files"""
    image_path = IMAGE_DIR / image_name
    if not image_path.exists():
        raise HTTPException(status_code=404, detail="Image not found")
    return FileResponse(image_path)


# ========================== CONVERSION ENDPOINTS ==========================

@app.post("/api/convert-slides")
async def convert_slides(data: dict):
    """Convert HTML slides to absolute positioning"""
    if conversion_service is None:
        raise HTTPException(status_code=503, detail="Converter not available")
    
    project_name = data.get("project_name")
    if not project_name:
        raise HTTPException(status_code=400, detail="Missing project_name")
    
    output_folder = OUTPUT_DIR / project_name
    if not output_folder.exists():
        raise HTTPException(status_code=404, detail="Project not found")
    
    html_folder = output_folder / "html"
    converted_folder = output_folder / "converted"
    
    if not html_folder.exists():
        raise HTTPException(status_code=404, detail="HTML folder not found")
    
    try:
        converted_folder.mkdir(exist_ok=True)
        conversion_result = conversion_service.convert_slides_batch(html_folder, converted_folder)
        
        return JSONResponse(content={
            "success": True,
            "conversion_result": conversion_result,
            "message": f"Converted {len(conversion_result['success'])} slides successfully"
        })
    except Exception as e:
        logger.error(f"Conversion error: {e}")
        raise HTTPException(status_code=500, detail=f"Conversion failed: {str(e)}")


@app.get("/api/conversion-status/{project_name}")
async def get_conversion_status(project_name: str):
    """Get conversion status for a project"""
    if conversion_service is None:
        raise HTTPException(status_code=503, detail="Converter not available")
    
    output_folder = OUTPUT_DIR / project_name
    if not output_folder.exists():
        raise HTTPException(status_code=404, detail="Project not found")
    
    html_folder = output_folder / "html"
    converted_folder = output_folder / "converted"
    
    if not html_folder.exists():
        raise HTTPException(status_code=404, detail="HTML folder not found")
    
    status = conversion_service.get_conversion_status(html_folder, converted_folder)
    
    return JSONResponse(content={
        "project_name": project_name,
        "status": status,
        "has_converted": converted_folder.exists()
    })


@app.get("/api/converter-health")
async def check_converter_health():
    """Check converter system health"""
    if conversion_service is None:
        return JSONResponse(content={
            "healthy": False,
            "message": "Converter is disabled or not initialized"
        })
    
    return JSONResponse(content=conversion_service.check_health())


# ========================== CAPTURE ENDPOINTS ==========================

@app.post("/api/capture-slides-batch")
async def capture_slides_batch(request: CaptureMultipleSlidesRequest):
    """Capture multiple slides and return as ZIP"""
    if capture_service is None or not capture_service.is_available():
        raise HTTPException(status_code=503, detail="Capture service not available")
    
    try:
        result = await capture_service.capture_multiple_slides(request.slides)
        return JSONResponse(content=result)
    except Exception as e:
        logger.error(f"Batch capture error: {e}")
        return JSONResponse(
            content={'success': False, 'error': str(e)},
            status_code=500
        )


@app.get("/api/download-capture/{filename}")
async def download_capture(filename: str):
    """Download captured slide or ZIP file"""
    try:
        if capture_service is None:
            raise HTTPException(status_code=503, detail="Capture service not available")
        
        file_path = capture_service.get_capture_file(filename)
        return FileResponse(
            file_path,
            filename=filename,
            media_type='application/octet-stream'
        )
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found")
    except Exception as e:
        logger.error(f"Download error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ========================== PDF EXPORT ENDPOINTS ==========================

@app.post("/api/export-slide-pdf")
async def export_slide_pdf(request: CaptureSlideRequest):
    """Export single slide to PDF"""
    if capture_service is None or not capture_service.is_available():
        raise HTTPException(status_code=503, detail="Capture service not available")
    
    try:
        result = await capture_service.export_slide_to_pdf(
            request.html_content,
            request.slide_name
        )
        return JSONResponse(content=result)
    except Exception as e:
        logger.error(f"PDF export error: {e}")
        return JSONResponse(
            content={'success': False, 'error': str(e)},
            status_code=500
        )


@app.post("/api/export-slides-pdf")
async def export_slides_pdf(request: CaptureMultipleSlidesRequest):
    """Export multiple slides to a single PDF"""
    if capture_service is None or not capture_service.is_available():
        raise HTTPException(status_code=503, detail="Capture service not available")
    
    try:
        result = await capture_service.export_slides_to_pdf(request.slides)
        return JSONResponse(content=result)
    except Exception as e:
        logger.error(f"PDF export error: {e}")
        return JSONResponse(
            content={'success': False, 'error': str(e)},
            status_code=500
        )


# ========================== HEALTH CHECK ==========================

@app.get("/api/health")
async def health_check():
    """Overall system health check"""
    return JSONResponse(content={
        "status": "healthy",
        "services": {
            "conversion": conversion_service is not None and conversion_service.is_available(),
            "capture": capture_service is not None and capture_service.is_available(),
            "workflow": workflow_service is not None and workflow_service.is_available(),
        }
    })


@app.get("/api/workflow-health")
async def check_workflow_health():
    """Check AI workflow system health"""
    try:
        if workflow_service is None:
            return JSONResponse(content={
                "healthy": False,
                "message": "Workflow service not initialized"
            })
        
        health = workflow_service.check_health()
        all_ok = health.get("workflow_available", False) and health.get("can_import", False)
        
        return JSONResponse(content={
            "healthy": all_ok,
            "details": health,
            "message": "Workflow system is ready" if all_ok else "Workflow system has issues"
        })
        
    except Exception as e:
        logger.error(f"Workflow health check error: {e}")
        return JSONResponse(content={
            "healthy": False,
            "error": str(e),
            "message": "Workflow health check failed"
        })


# ========================== TEMP FILE SERVING ==========================

@app.get("/api/temp-file")
async def serve_temp_file(path: str):
    """Serve file from temporary workflow directory via query parameter"""
    logger.info(f"📂 Serving temp file request: {path}")
    
    # Security: only allow serving from slidegen_workflow_*, slidegen_idea_*, or slidegen_manual_* directories
    normalized_path = path.replace('\\', '/')
    is_linux_workflow = normalized_path.startswith('/tmp/slidegen_workflow_')
    is_linux_idea = normalized_path.startswith('/tmp/slidegen_idea_')
    is_linux_manual = normalized_path.startswith('/tmp/slidegen_manual_')
    is_windows_temp = ('slidegen_workflow_' in normalized_path or 'slidegen_idea_' in normalized_path or 'slidegen_manual_' in normalized_path) and ('Temp' in path or 'temp' in path)
    
    if not (is_linux_workflow or is_linux_idea or is_linux_manual or is_windows_temp):
        logger.warning(f"⚠️ Access denied for path: {path}")
        raise HTTPException(status_code=403, detail="Access denied")
    
    file_path = Path(path)
    if not file_path.exists():
        logger.warning(f"⚠️ File not found: {path}")
        raise HTTPException(status_code=404, detail="File not found")
    
    logger.info(f"✅ Serving temp file: {path}")
    return FileResponse(file_path)


# ========================== STRUCTURE EDITOR ENDPOINTS ==========================

@app.get("/api/structure/current")
async def get_current_structure(session_id: str):
    """Get current document structure and text for editing"""
    logger.info(f"📋 Getting structure for session: {session_id}")
    
    # Check if session exists
    if session_id not in structure_sessions:
        raise HTTPException(status_code=404, detail="Session not found")
    
    session_data = structure_sessions[session_id]
    
    return JSONResponse(content={
        "document_title": session_data.get("document_title", "Untitled"),
        "document_text": session_data.get("document_text", ""),  # Full plain text for selection
        "document_structure": session_data.get("document_structure", []),
        "can_edit": True
    })


@app.put("/api/structure/update")
async def update_structure(request: StructureUpdateRequest):
    """Update document structure in session"""
    logger.info(f"💾 Updating structure for session: {request.session_id}")
    
    # Update session in memory
    if request.session_id not in structure_sessions:
        structure_sessions[request.session_id] = {}
    
    structure_sessions[request.session_id]["document_structure"] = request.document_structure
    
    # EXTRACT document_title from intro_slide
    document_title = "Untitled Document"
    if request.intro_slide:
        document_title = request.intro_slide.content or "Untitled Document"
        logger.info(f"📌 Document title: {document_title}")
    
    # PREPEND intro_slide if exists
    raw_structure = []
    if request.intro_slide:
        raw_structure.append(request.intro_slide.dict())
    
    # Add document structure
    raw_structure.extend([node.dict() for node in request.document_structure])
    
    # MERGE title + content nodes BEFORE saving
    def merge_title_content_nodes(nodes):
        """Merge consecutive title/content node pairs into single nodes"""
        merged = []
        i = 0
        
        while i < len(nodes):
            node = nodes[i]
            node_dict = node if isinstance(node, dict) else node
            
            # Check if this is a title node (has title, empty/missing content)
            has_title = node_dict.get('title', '').strip() != ''
            has_no_content = not node_dict.get('content', '').strip()
            
            # Check if next node is content node (has content, empty/missing title)
            is_next_content = False
            if i + 1 < len(nodes):
                next_node = nodes[i + 1]
                next_dict = next_node if isinstance(next_node, dict) else next_node
                has_no_title = not next_dict.get('title', '').strip()
                has_content = next_dict.get('content', '').strip() != ''
                same_level = next_dict.get('level') == node_dict.get('level')
                is_next_content = has_no_title and has_content and same_level
            
            # Merge if current is title and next is content
            if has_title and has_no_content and is_next_content:
                next_node = nodes[i + 1]
                next_dict = next_node if isinstance(next_node, dict) else next_node
                merged_node = {
                    'level': node_dict.get('level'),
                    'title': node_dict.get('title'),
                    'content': next_dict.get('content'),
                    'children': merge_title_content_nodes(node_dict.get('children', [])),
                    'is_transition': node_dict.get('is_transition', False)
                }
                merged.append(merged_node)
                i += 2  # Skip next node as it's been merged
            else:
                # Keep node as-is, but recurse into children
                merged_node = {
                    'level': node_dict.get('level'),
                    'title': node_dict.get('title', ''),
                    'content': node_dict.get('content', ''),
                    'children': merge_title_content_nodes(node_dict.get('children', [])),
                    'is_transition': node_dict.get('is_transition', False)
                }
                merged.append(merged_node)
                i += 1
        
        return merged
    
    # SAVE TO FILE with merged format
    try:
        # Ensure TEMP_DIR exists
        TEMP_DIR.mkdir(parents=True, exist_ok=True)
        
        structure_file = TEMP_DIR / f"structure_{request.session_id}.json"
        
        # Merge title/content pairs
        merged_structure = merge_title_content_nodes(raw_structure)
        
        structure_data = {
            "session_id": request.session_id,
            "document_title": document_title,  # Save title here!
            "document_structure": merged_structure
        }
        
        # Save to file
        with open(structure_file, 'w', encoding='utf-8') as f:
            json.dump(structure_data, f, indent=2, ensure_ascii=False)
        
        logger.info(f"✅ Structure updated: {request.session_id} ({len(raw_structure)} nodes → {len(merged_structure)} merged)")
        logger.info(f"📁 Saved to file: {structure_file}")
        
    except Exception as e:
        logger.error(f"Error saving structure to file: {e}")
        import traceback
        logger.error(traceback.format_exc())
    
    return JSONResponse(content={"success": True})


@app.post("/api/structure/approve")
async def approve_structure(request: StructureApproveRequest):
    """
    Approve structure and generate slide chunks using chunk_organizer logic.
    Returns State-compatible output for downstream slide generation.
    """
    logger.info(f"✅ Approving structure for session: {request.session_id}")
    
    THRESHOLD_CHAR_LIMIT = 1000
    
    try:
        # ======================== Load Structure ========================
        structure_file = TEMP_DIR / f"structure_{request.session_id}.json"
        
        if not structure_file.exists():
            logger.warning(f"No structure found for session: {request.session_id}")
            return JSONResponse(content={
                "success": False,
                "message": "No structure found for this session"
            }, status_code=404)
        
        with open(structure_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        document_structure = data.get('document_structure', [])
        document_title = data.get('document_title', 'Untitled Document')
        
        logger.info(f"📋 Loaded structure: {len(document_structure)} nodes")
        logger.info(f"📄 Document title: {document_title}")
        
        # ======================== Extract Special Slides ========================
       
        # Extract intro slide from level 0
        intro_slide = None
        intro_node = None
        filtered_structure = []
        
        for node in document_structure:
            if node.get('level') == 0:
                intro_node = node
            else:
                filtered_structure.append(node)
        
        if intro_node:
            intro_slide = f"INTRO SLIDE\n\nTitle: {intro_node.get('content', '')}"
            logger.info(f"✅ Extracted intro slide from level 0: {intro_node.get('content')}")
        # Manual mode: Don't auto-add intro unless user explicitly creates level 0 node
        # else:
        #     intro_slide = f"INTRO SLIDE\n\nTitle: {document_title}"
        #     logger.info(f"✅ Extracted intro slide from level 0: {document_title}")
        
        
        # Extract thank you slide (hybrid approach)
        thank_you_slide = None
        thank_you_node = None
        
        if filtered_structure and len(filtered_structure) > 0:
            last_node = filtered_structure[-1]
            # Check if last node title contains any thank you keyword
            thank_you_keywords = ['thank', 'thanks', 'thank you', 'cảm ơn', 'cam on']
            if any(keyword in last_node.get('title', '').lower() for keyword in thank_you_keywords):
                thank_you_node = last_node
                filtered_structure = filtered_structure[:-1]  # Remove from structure
                thank_you_slide = f"THANK YOU SLIDE\n\n{thank_you_node.get('title', '')}\n\n{thank_you_node.get('content', '')}"
                logger.info(f"✅ Extracted thank you slide from last node: {thank_you_node.get('title')}")
        
        # Manual mode: Don't auto-add thank you unless user explicitly created one
        # Fallback disabled - user has full control
        # if not thank_you_slide:
        #     thank_you_slide = "THANK YOU SLIDE\n\nThank you for your attention!"
        #     logger.info(f"📝 Using default thank you slide")
        
        document_structure = filtered_structure
        
        # ======================== Renumber Structure ========================
        
        def _add_numbering_recursively(nodes, counters=None, parent_level=0):
            counters = [0] * 6 if counters is None else counters
            for node in nodes:
                # For transition nodes: peek at next index (increment then decrement)
                if node.get('is_transition') == True:
                    # Temporarily increment to get the "next" index
                    counters[parent_level] += 1
                    for i in range(parent_level + 1, len(counters)):
                        counters[i] = 0
                    
                    # Assign index
                    index = ".".join(map(str, counters[:parent_level + 1]))
                    node['index'] = index
                    node['numbered_title'] = f"{index}. {node['title']}"
                    node['level'] = parent_level + 1
                    
                    # Decrement back (so next content node gets same index)
                    counters[parent_level] -= 1
                else:
                    # For content nodes: normal increment
                    counters[parent_level] += 1
                    for i in range(parent_level + 1, len(counters)):
                        counters[i] = 0
                    
                    # Assign index
                    index = ".".join(map(str, counters[:parent_level + 1]))
                    node['index'] = index
                    node['numbered_title'] = f"{index}. {node['title']}"
                    node['level'] = parent_level + 1
                
                if node.get('children'):
                    _add_numbering_recursively(node['children'], list(counters), parent_level + 1)
            return nodes
        
        renumbered_structure = _add_numbering_recursively(document_structure)
        logger.info(f"✅ Renumbered structure")
        
        # ======================== Generate Slides ========================
        
        final_slides_content = []
        final_slides_positions = []
        final_slides_context_maps = []
        final_slides_types = []
        
        # Add intro slide first
        if intro_slide:
            final_slides_content.append(intro_slide)
            final_slides_positions.append(['0'])
            final_slides_context_maps.append("INTRO SLIDE")
            final_slides_types.append('intro')
            logger.info("✅ Added intro slide")
        
        # Helper functions
        def _create_context_map(outline_nodes, active_indices):
            map_lines = []
            # Ensure active_indices is a list
            if not isinstance(active_indices, list):
                active_indices = [active_indices]
            
            def _generate(nodes, indent):
                for node in nodes:
                    node_index = node.get("index", "")
                    # Check if this node's index is in active_indices
                    marker = " <--- YOU ARE HERE" if node_index in active_indices else ""
                    map_lines.append(f'{"  " * indent}{node_index}. {node.get("title", "")}{marker}')
                    if node.get("children"):
                        _generate(node["children"], indent + 1)
            _generate(outline_nodes, 0)
            return "\n".join(map_lines)
        
        def _render_slide_content(group_header, items):
            parts = [group_header] if group_header else []
            if len(items) > 1:
                for item in items:
                    parts.append(f"\n\t- {item.get('title', '')}")
                    if item.get('content'):
                        indented = "\n\t  ".join(item['content'].splitlines())
                        parts.append(f"\t  {indented}")
            else:
                item = items[0]
                parts.append(f"\n{item.get('numbered_title', item.get('title', ''))}")
                if item.get('content'):
                    parts.append(f"{item['content']}")
            return "\n".join(parts)
        
        def _create_slides_from_leaves(parent, leaves):
            if not leaves:
                return
            buffer = []
            # Include parent's title AND content in header
            header_parts = [parent.get('numbered_title', '')]
            if parent.get('content'):
                header_parts.append(parent['content'])
            header = '\n'.join(header_parts)
            
            for item in leaves:
                tentative_buffer = buffer + [item]
                if len(_render_slide_content(header, tentative_buffer)) > THRESHOLD_CHAR_LIMIT and buffer:
                    content = _render_slide_content(header, buffer)
                    pos = [i.get('index', '') for i in buffer]
                    final_slides_content.append(content)
                    final_slides_positions.append(pos)
                    final_slides_context_maps.append(_create_context_map(renumbered_structure, pos))
                    final_slides_types.append('content')
                    buffer = [item]
                else:
                    buffer.append(item)
            if buffer:
                content = _render_slide_content(header, buffer)
                pos = [i.get('index', '') for i in buffer]
                final_slides_content.append(content)
                final_slides_positions.append(pos)
                final_slides_context_maps.append(_create_context_map(renumbered_structure, pos))
                final_slides_types.append('content')
        
        def _traverse_dfs(nodes, is_root_level=False):
            for node in nodes:
                try:
                    # Create transition slide ONLY if is_transition=true
                    if is_root_level and node.get('is_transition') == True:
                        transition_content = f"TRANSITION SLIDE\n\n{node.get('numbered_title', node.get('title', ''))}"
                        pos = [f"{node.get('index', '')}.T"]
                        final_slides_content.append(transition_content)
                        final_slides_positions.append(pos)
                        
                        node_index_list = [node.get('index', '')]
                        logger.info(f"🔍 Creating transition context map with indices: {node_index_list} (type: {type(node_index_list)})")
                        context_map = _create_context_map(renumbered_structure, node_index_list)
                        final_slides_context_maps.append(context_map)
                        final_slides_types.append('transition')
                        logger.info(f"✅ Created transition slide: {node.get('numbered_title', '')}")
                    
                    
                    # Check if node has children
                    has_leaves = any(not child.get("children") for child in node.get("children", []))
                    has_sub_trees = any(child.get("children") for child in node.get("children", []))
                    
                    # Process content ONLY for non-transition nodes WITHOUT children
                    # If node has children, content will be merged via _create_slides_from_leaves
                    # Also create slide if node has title but no content (use title as content)
                    node_has_content = node.get("content") or node.get("title")
                    if node_has_content and node.get('is_transition') != True and not has_leaves and not has_sub_trees:
                        pos = [node.get('index', '')]
                        # If content is empty but title exists, use title as content
                        if not node.get("content") and node.get("title"):
                            node_with_content = dict(node)
                            node_with_content["content"] = node.get("title", "")
                            final_slides_content.append(_render_slide_content("", [node_with_content]))
                        else:
                            final_slides_content.append(_render_slide_content("", [node]))
                        final_slides_positions.append(pos)
                        
                        logger.info(f"🔍 Creating content context map with indices: {pos} (type: {type(pos)})")
                        context_map = _create_context_map(renumbered_structure, pos)
                        final_slides_context_maps.append(context_map)
                        final_slides_types.append('content')
                    
                    # Process children
                    leaves = [child for child in node.get("children", []) if not child.get("children")]
                    if leaves:
                        _create_slides_from_leaves(node, leaves)
                    
                    sub_trees = [child for child in node.get("children", []) if child.get("children")]
                    if sub_trees:
                        _traverse_dfs(sub_trees, is_root_level=False)
                        
                except Exception as e:
                    logger.error(f"❌ Error in _traverse_dfs for node {node.get('title')}: {e}")
                    logger.error(f"   Node data: {node}")
                    import traceback
                    logger.error(traceback.format_exc())
                    raise
        
        # Traverse structure and create content slides
        _traverse_dfs(renumbered_structure, is_root_level=True)
        
        # Add thank you slide at the end
        if thank_you_slide:
            final_slides_content.append(thank_you_slide)
            final_slides_positions.append(['999'])
            final_slides_context_maps.append("THANK YOU SLIDE")
            final_slides_types.append('thank_you')
            logger.info("✅ Added thank you slide")
        
        logger.info(f"✅ Generated {len(final_slides_content)} slides:")
        logger.info(f"   - Intro: {final_slides_types.count('intro')}")
        logger.info(f"   - Transitions: {final_slides_types.count('transition')}")
        logger.info(f"   - Content: {final_slides_types.count('content')}")
        logger.info(f"   - Thank You: {final_slides_types.count('thank_you')}")
        
        # ======================== Sort Slides ========================
        
        def _get_version_key(pos_str):
            # Type safety: ensure pos_str is a string
            if isinstance(pos_str, (list, tuple)):
                logger.warning(f"⚠️ _get_version_key received {type(pos_str)}: {pos_str}, converting to string")
                pos_str = str(pos_str[0]) if pos_str else '0'
            
            pos_str = str(pos_str)  # Force to string
            
            if pos_str == '0':
                return (0,)
            if pos_str == '999':
                return (999999,)
            if pos_str.endswith('.T'):
                # Transition slides: place before content of same section
                parts = pos_str[:-2].split('.')
                return tuple(int(p) for p in parts) + (0,)  # e.g. "1.T" → (1, 0)
            
            # Content slides: place after transitions
            parts = pos_str.split('.')
            return tuple(int(p) for p in parts) + (1,)  # e.g. "1" → (1, 1)
        
        combined = list(zip(final_slides_content, final_slides_positions, final_slides_context_maps, final_slides_types))
        
        def sort_key(slide_data):
            return _get_version_key(slide_data[1][0]) if slide_data[1] else ()
        
        sorted_combined = sorted(combined, key=sort_key)
        
        if sorted_combined:
            sorted_chunks, sorted_positions, sorted_maps, sorted_types = zip(*sorted_combined)
        else:
            sorted_chunks, sorted_positions, sorted_maps, sorted_types = [], [], [], []
        
        logger.info(f"✅ Sorted {len(sorted_chunks)} slides")
        
        # Log samples
        for i, (chunk, slide_type) in enumerate(list(zip(sorted_chunks, sorted_types))[:3]):
            logger.info(f"\n--- Sample Slide {i+1} ({slide_type}) ---\n{chunk[:200]}...\n")
        
        # ======================== Save Results to File ========================
        
        # CONVERT slide_chunks to workflow-compatible format
        # slide_planner expects dicts with: position, heading, text_content
        converted_chunks = []
        for i, (chunk_text, chunk_type, positions) in enumerate(zip(sorted_chunks, sorted_types, sorted_positions)):
            # Determine position value based on type
            if chunk_type == 'intro':
                position = 0
            elif chunk_type == 'transition':
                position = -1
            elif chunk_type == 'thank_you':
                position = -2
            else:
                position = i  # Content slides get sequential position
            
            # Extract heading from chunk text (first line usually)
            lines = chunk_text.strip().split('\n')
            if chunk_type == 'intro':
                heading = 'INTRO'
            elif chunk_type == 'transition':
                heading = 'TRANSITION'
            elif chunk_type == 'thank_you':
                heading = 'THANK_YOU'
            else:
                # For content, extract heading from first line
                heading = lines[0] if lines else ''
                # Remove numbering if present (e.g. "1.2. Title" -> "Title")
                import re
                heading = re.sub(r'^\d+(\.\d+)*\.?\s*', '', heading)
            
            # Create workflow-compatible chunk dict
            chunk_dict = {
                'position': position,
                'heading': heading,
                'text_content': chunk_text
            }
            converted_chunks.append(chunk_dict)
        
        logger.info(f"✅ Converted {len(converted_chunks)} chunks to workflow format")
        
        # Log sample converted chunk
        if converted_chunks:
            logger.info(f"Sample converted chunk: {converted_chunks[0]}")
        
        try:
            # Save approved result to file for debugging/monitoring
            approved_file = TEMP_DIR / f"approved_structure_{request.session_id}.json"
            
            result_data = {
                "success": True,
                "message": "Structure approved successfully",
                "session_id": request.session_id,
                "document_title": document_title,
                "document_outline": renumbered_structure,
                "slide_chunks": converted_chunks,  # Use converted format
                "slide_positions": list(sorted_positions),
                "slide_context_maps": list(sorted_maps),
                "slide_types": list(sorted_types)
            }
            
            with open(approved_file, 'w', encoding='utf-8') as f:
                json.dump(result_data, f, indent=2, ensure_ascii=False)
            
            logger.info(f"💾 Saved approved result to: {approved_file}")
            
        except Exception as e:
            logger.error(f"⚠️ Error saving approved result to file: {e}")
        
        # ======================== Return State Output ========================
        
        return JSONResponse(content={
            "success": True,
            "message": "Structure approved successfully",
            "session_id": request.session_id,
            "document_title": document_title,
            "document_outline": renumbered_structure,
            "slide_chunks": converted_chunks,  # Return converted format
            "slide_positions": list(sorted_positions),
            "slide_context_maps": list(sorted_maps),
            "slide_types": list(sorted_types)
        })
        
    except Exception as e:
        logger.error(f"Error approving structure: {e}")
        import traceback
        logger.error(traceback.format_exc())
        return JSONResponse(content={
            "success": False,
            "message": f"Error processing structure: {str(e)}"
        }, status_code=500)


@app.post("/api/structure/test-session")
async def create_test_session():
    """Create test session with sample structure for testing"""
    import uuid
    session_id = str(uuid.uuid4())
    
    # Sample structure
    structure_sessions[session_id] = {
        "document_title": "Test Document - Manual Structure Editor",
        "document_structure": [
            {
                "level": 1,
                "title": "Giới thiệu",
                "content": "Đây là phần giới thiệu của tài liệu. Nội dung này khá dài để test tính năng split. " * 50,
                "children": [
                    {
                        "level": 2,
                        "title": "Bối cảnh",
                        "content": "Nội dung về bối cảnh dự án.",
                        "children": []
                    },
                    {
                        "level": 2,
                        "title": "Mục tiêu",
                        "content": "Mục tiêu của dự án là xây dựng hệ thống tạo slide tự động.",
                        "children": []
                    }
                ]
            },
            {
                "level": 1,
                "title": "Phương pháp",
                "content": "Mô tả các phương pháp được sử dụng trong dự án.",
                "children": [
                    {
                        "level": 2,
                        "title": "Công nghệ",
                        "content": "Sử dụng Python, FastAPI, React.",
                        "children": []
                    }
                ]
            },
            {
                "level": 1,
                "title": "Kết luận",
                "content": "Tóm tắt các kết quả đạt được.",
                "children": []
            }
        ]
    }
    
    logger.info(f"✅ Created test session: {session_id}")
    
    return JSONResponse(content={
        "success": True,
        "session_id": session_id,
        "redirect_url": f"/frontend/structure-editor.html?session={session_id}"
    })


@app.post("/api/cleanup-temp")
async def cleanup_temp_directory(request: Request):
    """Clean up temporary workflow directory when user navigates to editor"""
    try:
        data = await request.json()
        temp_dir = data.get('temp_dir')
        
        if not temp_dir:
            return JSONResponse(content={"success": False, "message": "No temp directory specified"})
        
        # Security: only allow deleting slidegen_workflow_* or slidegen_idea_* directories
        normalized_path = temp_dir.replace('\\', '/')
        is_linux_workflow = normalized_path.startswith('/tmp/slidegen_workflow_')
        is_linux_idea = normalized_path.startswith('/tmp/slidegen_idea_')
        is_windows_temp = ('slidegen_workflow_' in normalized_path or 'slidegen_idea_' in normalized_path) and ('Temp' in temp_dir or 'temp' in temp_dir)
        
        if not (is_linux_workflow or is_linux_idea or is_windows_temp):
            raise HTTPException(status_code=403, detail="Invalid temp directory")
        
        temp_path = Path(temp_dir)
        if temp_path.exists():
            shutil.rmtree(temp_path)
            logger.info(f"🗑️ Cleaned up temp directory: {temp_dir}")
            return JSONResponse(content={"success": True, "message": "Temp directory cleaned"})
        else:
            return JSONResponse(content={"success": False, "message": "Temp directory not found"})
    
    except Exception as e:
        logger.error(f"Cleanup error: {e}")
        return JSONResponse(content={"success": False, "error": str(e)}, status_code=500)


@app.post("/api/generate-from-idea")
async def generate_from_idea(
    request: Request,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """Generate slides from natural language idea using LLM (VIP Required)"""
    # Check VIP status
    if not current_user.is_vip:
        raise HTTPException(
            status_code=403,
            detail="VIP membership required to use the 'Generate from Idea' feature. Please upgrade to VIP."
        )
    
    try:
        data = await request.json()
        title = data.get('title', '').strip()
        prompt = data.get('prompt', '').strip()
        content_length = data.get('content_length', 'standard')
        enable_transition_slides = data.get('enable_transition_slides', True)
        enable_thank_you_slide = data.get('enable_thank_you_slide', True)
        
        if not title or not prompt:
            raise HTTPException(status_code=400, detail="Title and prompt are required")
        
        # Map content length to EXACT slide count (TOTAL slides including all types)
        # User selects a preset, we give them a specific target total
        length_map = {
            'brief': {'total': 6, 'desc': 'Quick overview'},
            'standard': {'total': 10, 'desc': 'Balanced coverage'},
            'detailed': {'total': 15, 'desc': 'Comprehensive'},
            'extensive': {'total': 22, 'desc': 'In-depth analysis'}
        }
        length_info = length_map.get(content_length, length_map['standard'])
        
        # EXACT calculation - NO RANGES
        target_total = length_info['total']
        
        # Fixed counts
        intro_count = 1  # Always exactly 1 intro
        thank_you_count = 1 if enable_thank_you_slide else 0
        
        # Calculate sections and transitions based on target
        # Rule: ~3-4 content slides per section
        if enable_transition_slides:
            # Estimate sections: for small presentations 2-3 sections, for large 4-5
            if target_total <= 8:
                num_sections = 2
            elif target_total <= 12:
                num_sections = 3
            elif target_total <= 18:
                num_sections = 4
            else:
                num_sections = 5
            transition_count = num_sections - 1  # Transition between sections (not before first)
        else:
            num_sections = 3  # Default for organization
            transition_count = 0
        
        # Calculate exact content slides
        content_count = target_total - intro_count - transition_count - thank_you_count
        content_count = max(2, content_count)  # At least 2 content slides
        
        # Recalculate actual total after adjustments
        actual_total = intro_count + content_count + transition_count + thank_you_count
        
        logger.info(f"🤖 Generating slides from idea:")
        logger.info(f"   Title: {title}")
        logger.info(f"   Prompt length: {len(prompt)} chars")
        logger.info(f"   🎯 Target EXACT: {target_total} slides")
        logger.info(f"   📊 Breakdown: intro={intro_count}, content={content_count}, transitions={transition_count}, thank_you={thank_you_count}")
        logger.info(f"   📊 Actual total: {actual_total} slides")
        logger.info(f"   📊 Sections: {num_sections}")
        logger.info(f"   Options: transitions={enable_transition_slides}, thank_you={enable_thank_you_slide}")
        
        # Call LLM to generate slide structure
        from src.models.llm import LLM
        import os
        
        gemini_api_key = os.getenv('GEMINI_API_KEY')
        if not gemini_api_key:
            raise HTTPException(status_code=500, detail="GEMINI_API_KEY not configured")
        
        # Import config for idea structure LLM
        from src.utils.config import (
            IDEA_STRUCTURE_LLM_MODEL,
            IDEA_STRUCTURE_LLM_TEMPERATURE,
            IDEA_STRUCTURE_LLM_MAX_TOKENS,
            LLM_BASE_URL
        )
        
        llm = LLM(
            api_keys=[gemini_api_key],
            base_url=LLM_BASE_URL,
            model=IDEA_STRUCTURE_LLM_MODEL,
            temperature=IDEA_STRUCTURE_LLM_TEMPERATURE,
            max_tokens=IDEA_STRUCTURE_LLM_MAX_TOKENS
        )
        
        # Build structured prompt for slide generation with EXACT counts
        structure_prompt = f"""You are a presentation structure generator. Given a title and description, create a logical slide structure.

PRESENTATION TITLE: {title}

USER DESCRIPTION:
{prompt}

CONTENT DEPTH: {content_length.upper()} ({length_info['desc']})

===== EXACT SLIDE COUNT REQUIREMENTS =====
You MUST generate EXACTLY this many slides:
- INTRO slides: {intro_count} (position 0)
- CONTENT slides: {content_count} (position 1, 2, 3... up to {content_count})
- TRANSITION slides: {transition_count} (position -1)
- THANK YOU slides: {thank_you_count} (position -2)
- TOTAL: {actual_total} slides EXACTLY

DO NOT generate more or fewer slides than specified above.
==========================================

ORGANIZATION:
- Organize content into {num_sections} major sections
- Distribute the {content_count} content slides across these sections

INSTRUCTIONS:
1. First slide MUST be intro/title slide (position 0)
2. Generate EXACTLY {content_count} content slides (position 1, 2, 3... up to {content_count})
3. {"Add EXACTLY " + str(transition_count) + " transition slides (position -1) between major sections" if transition_count > 0 else "DO NOT add any transition slides"}
4. {"Add EXACTLY 1 thank you slide (position -2) at the end" if thank_you_count > 0 else "DO NOT add any thank you slide"}
5. **CRITICAL**: For each slide's text_content, provide COMPLETE and DETAILED information:
   - DO NOT just list bullet point headings like "- Introduction" or "- Benefits"
   - INSTEAD, write out FULL information with specific details, data, and explanations
   - Example GOOD: "- Machine learning improves efficiency by 80% through automated data processing"
   - Example BAD: "- Benefits" or "- Machine Learning"
   - Each bullet point should be 10-30 words with concrete information

OUTPUT FORMAT (JSON):
{{
  "slide_chunks": [
    {{"position": 0, "heading": "", "text_content": "Title: {title}\\nSubtitle: [subtitle]"}},
    {{"position": 1, "heading": "Section Title", "text_content": "- Detailed point 1\\n- Detailed point 2\\n- Detailed point 3"}},
    ... more content slides with position 2, 3, etc. up to {content_count} ...
    {'{{"position": -1, "heading": "TRANSITION", "text_content": "Moving to: Next Section"}},' if transition_count > 0 else ''}
    {'{{"position": -2, "heading": "THANK_YOU", "text_content": "Thank you for your attention!"}}' if thank_you_count > 0 else ''}
  ]
}}

RULES:
- position: 0 for intro, positive integers (1,2,3...) for content, -1 for transitions, -2 for thank you
- heading: Empty string for intro, "TRANSITION" for transitions, "THANK_YOU" for thank you, actual titles for content
- VERIFY your output has EXACTLY {actual_total} total slides before responding

Now generate the slide structure as valid JSON only (no markdown code blocks):"""
        
        # Call LLM
        logger.info("📤 Calling LLM for slide structure generation...")
        response_text = llm.invoke(structure_prompt)
        logger.info(f"📥 LLM response length: {len(response_text)} chars")
        
        # Parse JSON response
        import json
        import re
        
        # Try to extract JSON from response (handle markdown code blocks)
        json_match = re.search(r'\{[\s\S]*\}', response_text)
        if json_match:
            json_str = json_match.group(0)
        else:
            json_str = response_text
        
        try:
            structure_data = json.loads(json_str)
            slide_chunks = structure_data.get('slide_chunks', [])
        except json.JSONDecodeError as e:
            logger.error(f"Failed to parse LLM response as JSON: {e}")
            logger.error(f"Response text: {response_text[:500]}")
            raise HTTPException(status_code=500, detail="Failed to parse LLM response")
        
        if not slide_chunks:
            raise HTTPException(status_code=500, detail="No slides generated")
        
        # Validate and adjust slide count
        content_slides = [s for s in slide_chunks if isinstance(s.get('position'), int) and s.get('position') > 0]
        intro_slides = [s for s in slide_chunks if s.get('position') == 0]
        transition_slides_gen = [s for s in slide_chunks if s.get('position') == -1]
        thank_you_slides = [s for s in slide_chunks if s.get('position') == -2]
        
        total_slides = len(slide_chunks)
        
        logger.info(f"📊 Slide breakdown (LLM generated):")
        logger.info(f"   Intro: {len(intro_slides)} (expected: {intro_count})")
        logger.info(f"   Content: {len(content_slides)} (expected: {content_count})")
        logger.info(f"   Transitions: {len(transition_slides_gen)} (expected: {transition_count})")
        logger.info(f"   Thank you: {len(thank_you_slides)} (expected: {thank_you_count})")
        logger.info(f"   TOTAL: {total_slides} slides (expected: {actual_total})")
        
        # Validate slide counts match expectations
        if len(content_slides) != content_count:
            logger.warning(f"⚠️  Content slides mismatch: got {len(content_slides)}, expected {content_count}")
        
        if total_slides != actual_total:
            logger.warning(f"⚠️  Total slides mismatch: got {total_slides}, expected {actual_total}")
        else:
            logger.info(f"✅ Total slide count matches expected: {actual_total}")
        
        logger.info(f"✅ Generated {len(slide_chunks)} total slide chunks")
        
        # Save to temp file (mimic DOCX upload flow)
        import uuid
        session_id = str(uuid.uuid4())[:8]
        temp_filename = f"idea_{session_id}.json"
        temp_path = TEMP_DIR / temp_filename
        
        # Save slide chunks
        with open(temp_path, 'w', encoding='utf-8') as f:
            json.dump({
                'title': title,
                'slide_chunks': slide_chunks,
                'enable_transition_slides': enable_transition_slides,
                'enable_thank_you_slide': enable_thank_you_slide
            }, f, indent=2, ensure_ascii=False)
        
        # Save options file
        options_path = TEMP_DIR / f"{temp_filename}.options.json"
        with open(options_path, 'w') as f:
            json.dump({
                'enable_transition_slides': enable_transition_slides,
                'enable_thank_you_slide': enable_thank_you_slide
            }, f, indent=2)
        
        logger.info(f"💾 Saved slide structure to {temp_path}")
        
        return JSONResponse(content={
            'session_id': session_id,
            'filename': temp_filename,
            'slide_count': len(slide_chunks),
            'is_from_idea': True  # Flag to use different streaming endpoint
        })
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error generating from idea: {e}")
        import traceback
        logger.error(traceback.format_exc())
        raise HTTPException(status_code=500, detail=f"Failed to generate slides: {str(e)}")


@app.get("/api/generate-from-idea-stream")
async def generate_from_idea_stream(
    session_id: str = Query(..., description="Session ID for idea-based generation"),
    current_user: UserResponse = Depends(get_current_user_from_token_param)
):
    """Stream processing of idea-generated slides with Server-Sent Events - VIP Required"""
    
    logger.info(f"🔐 generate_from_idea_stream: Authenticated user: {current_user.email}")
    
    # Check VIP status
    if not current_user.is_vip:
        raise HTTPException(
            status_code=403,
            detail="VIP membership required to use the 'Generate from Idea' feature. Please upgrade to VIP."
        )
    
    
    # Validate session file
    temp_filename = f"idea_{session_id}.json"
    temp_path = TEMP_DIR / temp_filename
    if not temp_path.exists():
        raise HTTPException(status_code=400, detail="Invalid session")
    
    # Load slide chunks and options
    try:
        with open(temp_path, 'r', encoding='utf-8') as f:
            idea_data = json.load(f)
        
        title = idea_data.get('title', 'Untitled')
        slide_chunks = idea_data.get('slide_chunks', [])
        slide_options = {
            'enable_transition_slides': idea_data.get('enable_transition_slides', True),
            'enable_thank_you_slide': idea_data.get('enable_thank_you_slide', True)
        }
        
        logger.info(f"📋 Loaded idea data: {len(slide_chunks)} chunks, title='{title}'")
        
    except Exception as e:
        logger.error(f"Failed to load idea data: {e}")
        raise HTTPException(status_code=500, detail="Failed to load idea data")
    
    # Create project name
    project_name = f"idea_{session_id}"
    output_folder = OUTPUT_DIR / project_name
    
    # Handle existing project folder
    if output_folder.exists():
        try:
            shutil.rmtree(output_folder)
        except PermissionError:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            project_name = f"{project_name}_{timestamp}"
            output_folder = OUTPUT_DIR / project_name
    
    output_folder.mkdir(parents=True, exist_ok=True)
    
    # Store temp directory path for cleanup later
    temp_workflow_dir = None
    
    async def event_generator() -> AsyncGenerator[str, None]:
        """Generate Server-Sent Events for workflow progress"""
        nonlocal temp_workflow_dir
        
        # Queue for events
        event_queue = asyncio.Queue()
        workflow_done = asyncio.Event()
        workflow_error = None
        
        def stream_callback(event_type: str, data: dict):
            """Callback to receive events from workflow"""
            try:
                logger.info(f"🔔 Stream callback: {event_type} - {data}")
                
                # For slide_completed, keep the temp path for real-time access
                if event_type == 'slide_completed' and 'html_path' in data:
                    temp_path_slide = data['html_path']
                    data['html_path'] = temp_path_slide
                    data['is_temp'] = True
                    logger.info(f"🎯 Streaming temp file: {temp_path_slide}")
                
                # Send all events immediately for real-time streaming
                event_queue.put_nowait((event_type, data))
            except Exception as e:
                logger.error(f"Error in stream callback: {e}")
        
        async def run_workflow():
            """Run workflow in background with pre-loaded slide chunks"""
            nonlocal workflow_error, temp_workflow_dir
            try:
                if not workflow_service or not workflow_service.is_available():
                    raise Exception("Workflow service not available")
                
                # Call workflow with slide_chunks directly (skip document conversion)
                workflow_result = await workflow_service.process_idea_with_workflow(
                    slide_chunks=slide_chunks,
                    title=title,
                    output_folder=str(output_folder),
                    callback=stream_callback,
                    slide_options=slide_options
                )
                
                if not workflow_result:
                    raise Exception("Workflow did not produce valid output")
                
                # Unpack result
                result_zip, temp_workflow_dir = workflow_result
                logger.info(f"📁 Temp workflow directory: {temp_workflow_dir}")
                logger.info(f"📦 Result ZIP: {result_zip}")
                
                if result_zip and os.path.exists(result_zip):
                    # Extract ZIP to output folder
                    html_folder = output_folder / "html"
                    html_folder.mkdir(exist_ok=True)
                    
                    with zipfile.ZipFile(result_zip, 'r') as zip_ref:
                        for file_info in zip_ref.infolist():
                            if not file_info.filename.endswith('/') and 'temp_debug' not in file_info.filename:
                                try:
                                    zip_ref.extract(file_info, output_folder)
                                except Exception as e:
                                    logger.warning(f"Error extracting {file_info.filename}: {e}")
                    
                    # Convert slides if converter available
                    if conversion_service and conversion_service.is_available():
                        converted_folder = output_folder / "converted"
                        converted_folder.mkdir(exist_ok=True)
                        conversion_service.convert_slides_batch(html_folder, converted_folder)
                    
                    # Count total slides
                    html_files = list(html_folder.glob("*.html"))
                    total_slides = len(html_files)
                    logger.info(f"📁 Found {total_slides} HTML files")
                    
                    # Save to database
                    project_id = save_workflow_to_database(
                        user_id=current_user.id,
                        project_name=project_name,
                        source_type='idea',
                        source_data={
                            'title': title,
                            'prompt': idea_data.get('prompt', ''),
                            'description': f'Slides generated from idea: {title}'
                        },
                        slide_options=slide_options,
                        output_folder=output_folder
                    )
                    
                    # Emit complete event with temp directory path for cleanup
                    await event_queue.put(('complete', {
                        'project_name': project_name,
                        'total_slides': total_slides,
                        'temp_dir': temp_workflow_dir,
                        'project_id': project_id
                    }))
                else:
                    raise Exception("Workflow did not produce output files")
                
            except Exception as e:
                workflow_error = str(e)
                logger.error(f"Workflow error: {e}")
                import traceback
                logger.error(traceback.format_exc())
            finally:
                workflow_done.set()
        
        # Start workflow task
        workflow_task = asyncio.create_task(run_workflow())
        
        try:
            while not workflow_done.is_set() or not event_queue.empty():
                try:
                    # Wait for event with timeout
                    event_type, data = await asyncio.wait_for(
                        event_queue.get(),
                        timeout=1.0
                    )
                    
                    # Format SSE message
                    event_data = json.dumps(data)
                    yield f"event: {event_type}\ndata: {event_data}\n\n"
                    
                except asyncio.TimeoutError:
                    # Send keepalive
                    yield f": keepalive\n\n"
                    continue
            
            # Check for errors
            if workflow_error:
                error_data = json.dumps({"message": workflow_error})
                yield f"event: error_event\ndata: {error_data}\n\n"
            
        except Exception as e:
            logger.error(f"Error in event generator: {e}")
            error_data = json.dumps({"message": str(e)})
            yield f"event: error_event\ndata: {error_data}\n\n"
        finally:
            # Ensure workflow completes
            try:
                await workflow_task
            except:
                pass
    
    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )


# ============================================================================
# PIXABAY API ENDPOINTS
# ============================================================================

@app.get("/api/pixabay/search/images")
async def search_pixabay_images(
    q: str = Query(..., description="Search query"),
    image_type: str = Query("photo", description="Image type: all, photo, illustration, vector"),
    orientation: str = Query("horizontal", description="Orientation: all, horizontal, vertical"),
    per_page: int = Query(20, ge=3, le=200, description="Results per page"),
    page: int = Query(1, ge=1, description="Page number"),
    category: str = Query(None, description="Category filter"),
    colors: str = Query(None, description="Color filter"),
    order: str = Query("popular", description="Order: popular, latest")
):
    """
    Tìm kiếm hình ảnh từ Pixabay
    """
    try:
        result = PixabayTool.search_images(
            q=q,
            image_type=image_type,
            orientation=orientation,
            per_page=per_page,
            page=page,
            category=category,
            colors=colors,
            order=order
        )
        return JSONResponse(content=result)
    except Exception as e:
        logger.error(f"Error searching Pixabay images: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/pixabay/search/videos")
async def search_pixabay_videos(
    q: str = Query(..., description="Search query"),
    video_type: str = Query("all", description="Video type: all, film, animation"),
    per_page: int = Query(20, ge=3, le=200, description="Results per page"),
    page: int = Query(1, ge=1, description="Page number"),
    category: str = Query(None, description="Category filter"),
    order: str = Query("popular", description="Order: popular, latest")
):
    """
    Tìm kiếm video từ Pixabay
    """
    try:
        result = PixabayTool.search_videos(
            q=q,
            video_type=video_type,
            per_page=per_page,
            page=page,
            category=category,
            order=order
        )
        return JSONResponse(content=result)
    except Exception as e:
        logger.error(f"Error searching Pixabay videos: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/pixabay/best-image")
async def get_best_pixabay_image(
    q: str = Query(..., description="Search query"),
    orientation: str = Query("horizontal", description="Orientation: horizontal, vertical"),
    quality: str = Query("large", description="Quality: preview, webformat, large"),
    image_type: str = Query("photo", description="Image type: photo, illustration, vector")
):
    """
    Lấy URL của hình ảnh tốt nhất cho từ khóa
    """
    try:
        result = PixabayTool.search_and_get_best_image(
            query=q,
            orientation=orientation,
            quality=quality,
            image_type=image_type
        )
        if result:
            return JSONResponse(content=result)
        else:
            raise HTTPException(status_code=404, detail="No images found")
    except Exception as e:
        logger.error(f"Error getting best Pixabay image: {e}")
        raise HTTPException(status_code=500, detail=str(e))

# ============================================================================
# TRANSLATION API
# ============================================================================

@app.get("/api/translate")
async def translate_text(
    q: str = Query(..., description="Text to translate"),
    source: str = Query("vi", description="Source language code (default: vi)"),
    target: str = Query("en", description="Target language code (default: en)")
):
    """
    Translate text from one language to another
    
    - **q**: Text to translate
    - **source**: Source language (default: Vietnamese - vi)
    - **target**: Target language (default: English - en)
    
    Returns translated text with caching and multiple API fallbacks
    """
    try:
        translator = get_translator()
        translated = translator.translate(q, source=source, target=target)
        
        return JSONResponse(content={
            "original": q,
            "translated": translated,
            "source": source,
            "target": target
        })
    except Exception as e:
        logger.error(f"Translation error: {e}")
        # Return original text if translation fails
        return JSONResponse(content={
            "original": q,
            "translated": q,
            "source": source,
            "target": target,
            "error": str(e)
        })