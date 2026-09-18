# File: backend/api/v1/generate_router.py

import os
import json
import logging
import asyncio
import math
import re
import uuid
from pathlib import Path
from datetime import datetime
from backend.api.schemas.generate_schema import (
    DocxPrepareResponse, ManualSessionResponse
)
from backend.api.schemas.auth_schema import UserResponse
from fastapi import APIRouter, UploadFile, File, Form, Depends, HTTPException, Query
from fastapi.responses import JSONResponse, StreamingResponse

# Config & Domain
from backend.config import DOCX_MAX_CHARS, OUTPUT_DIR, TEMP_DIR, THRESHOLD_CHAR_LIMIT

# Dependencies
from backend.api.dependencies import (
    get_current_active_user,
    get_current_user_from_token_param,
    get_project_service,
    get_structure_service,
    get_slide_service,
    get_slidegen_workflow_service,
    get_legacy_docx_generation_service,
    get_subscription_service,
)
from backend.application.services.subscription_service import SubscriptionService
from backend.infrastructure.database.repositories.sqlite_user_repo import SqliteUserRepository
from backend.api.constants import EventTypes

# Services
from backend.application.services.project_service import ProjectService
from backend.application.services.structure_service import StructureService
from backend.application.services.slide_service import SlideService
from backend.application.services.slidegen_workflow_service import SlideGenWorkflowService
from backend.application.services.legacy_docx_generation_service import (
    LegacyDocxGenerationService,
)

router = APIRouter(prefix="/generate", tags=["GENERATE"])
logger = logging.getLogger("SlideGen-Generation")


def _new_temp_docx_path(filename: str) -> Path:
    """Create a collision-resistant path for an uploaded DOCX."""
    safe_filename = Path(filename or "upload.docx").name
    return TEMP_DIR / f"{uuid.uuid4().hex}_{safe_filename}"


def _safe_unlink(path: Path) -> None:
    """Remove a temporary file without failing the request if it is locked."""
    try:
        if path.exists():
            path.unlink()
    except PermissionError as exc:
        logger.warning("Could not remove locked temporary file %s: %s", path, exc)
    except OSError as exc:
        logger.warning("Could not remove temporary file %s: %s", path, exc)


def _cleanup_temp_dir():
    """Helper to clean temp files"""
    try:
        if not TEMP_DIR.exists(): return
        for old_file in TEMP_DIR.glob("*"):
            if old_file.is_file():
                _safe_unlink(old_file)
    except: pass


def _check_docx_slide_limit(temp_path, subscription_service, user):
    """Reject oversized DOCX files before preprocessing or slide planning."""
    try:
        from docx import Document

        doc = Document(temp_path)
        heading_count = 0
        content_chars = 0

        def visit_paragraph(paragraph) -> None:
            nonlocal content_chars, heading_count
            text = (paragraph.text or '').strip()
            content_chars += len(text)
            if paragraph.style.name.startswith('Heading'):
                heading_count += 1

        for paragraph in doc.paragraphs:
            visit_paragraph(paragraph)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        visit_paragraph(paragraph)

        if DOCX_MAX_CHARS > 0 and content_chars > DOCX_MAX_CHARS:
            raise HTTPException(
                status_code=413,
                detail=(
                    f"Tài liệu quá dài ({content_chars:,} ký tự). "
                    f"Hệ thống chỉ hỗ trợ tối đa {DOCX_MAX_CHARS:,} ký tự cho mỗi DOCX. "
                    "Vui lòng rút ngắn tài liệu rồi thử lại."
                ),
            )

        chars_per_slide = max(1, THRESHOLD_CHAR_LIMIT)
        content_slide_estimate = math.ceil(content_chars / chars_per_slide)
        slide_count = max(heading_count, content_slide_estimate, 1)
        subscription_service.check_feature_access(user, "docx_upload", slide_count=slide_count)

    except HTTPException:
        raise
    except Exception as e:
        logger.warning(f"Could not estimate slide count: {e}")
        # A failed estimate must not be converted into a fake 25-slide
        # violation. Let the workflow perform the real validation instead.
        subscription_service.check_feature_access(user, "docx_upload")

def prepare_output_folder(name: str):
    """Helper to create unique output folder"""
    raw_name = str(name or "Untitled")
    raw_name = re.sub(r"[\x00-\x1f\x7f]", " ", raw_name)
    raw_name = re.sub(r"[\\/:*?\"<>|]", " ", raw_name)
    safe_name = re.sub(r"\s+", "_", raw_name).strip("._")
    if not safe_name:
        safe_name = "Untitled"

    if len(safe_name) > 120:
        safe_name = safe_name[:120].rstrip("._")

    folder = OUTPUT_DIR / safe_name
    if folder.exists():
        folder = OUTPUT_DIR / f"{safe_name}_{int(datetime.now().timestamp())}"
    folder.mkdir(parents=True, exist_ok=True)
    return folder

def sanitize_event_data(event_type: str, data: dict):
    """Helper to mark temp files for frontend viewer"""
    new_data = data.copy()
    if event_type == 'slide_completed':
        if 'xml_path' in new_data: new_data['is_temp'] = True
        if 'image_path' in new_data: new_data['is_temp'] = True
    return new_data


def _workflow_error_payload(error: Exception) -> dict:
    """Convert workflow exceptions into a safe, useful SSE error payload."""
    status_code = getattr(error, "status_code", None)
    detail = getattr(error, "detail", None)
    message = detail if isinstance(detail, str) else str(error or "").strip()
    normalized = message.lower()

    if "does not contain heading structure" in normalized and "auto_heading" in normalized:
        message = (
            "Tài liệu không có tiêu đề (Heading) để chia slide. "
            "Vui lòng chọn chế độ ‘Tự động không tiêu đề’ hoặc thêm Heading vào tài liệu rồi thử lại."
        )
        status_code = status_code or 400
        code = "DOCUMENT_MISSING_HEADINGS"
    elif (
        "recursion limit" in normalized
        or "maximum number of steps" in normalized
        or "too many slides" in normalized
    ):
        message = (
            "Tài liệu hoặc số lượng slide quá lớn nên workflow không thể hoàn tất. "
            "Vui lòng rút ngắn tài liệu rồi thử lại."
        )
        status_code = status_code or 413
        code = "DOCUMENT_TOO_LONG"
    elif "tài liệu quá dài" in normalized or "exceeds the slide limit" in normalized:
        status_code = status_code or 413
        code = "DOCUMENT_TOO_LONG"
    elif not message:
        message = "Không thể tạo slide. Vui lòng thử lại."
        code = "WORKFLOW_ERROR"
    else:
        if not normalized.startswith(
            ("tài liệu", "không thể", "workflow", "gói", "phiên", "dịch vụ")
        ):
            message = f"Không thể tạo slide: {message}"
        code = "DOCUMENT_TOO_LONG" if status_code == 413 else "WORKFLOW_ERROR"

    return {
        "code": code,
        "message": message,
        **({"status_code": status_code} if status_code else {}),
    }


async def _run_stream_task(queue: asyncio.Queue, runner, label: str) -> None:
    """Run a background workflow and preserve its real error for the SSE client."""
    try:
        await runner()
    except Exception as exc:
        logger.exception("[%s] Workflow failed", label)
        await queue.put((EventTypes.ERROR_EVENT, _workflow_error_payload(exc)))


def count_workflow_slides(output_folder: Path, output_format: str = "xml") -> int:
    """Count generated XML artifacts."""
    folder_name = "xml"
    base_output = output_folder / "output"
    slide_dir = (base_output / folder_name) if (base_output / folder_name).exists() else (output_folder / folder_name)
    return len(list(slide_dir.glob("*.xml"))) if slide_dir.exists() else 0


# ==========================================
# 1. GENERATE FROM FILE (DOCX)
# ==========================================


@router.post("/docx")
async def generate_from_docx_normal(
    file: UploadFile = File(...),
    output_format: str = Form("xml"),
    current_user: UserResponse = Depends(get_current_active_user),
    project_service: ProjectService = Depends(get_project_service),
    slide_service: SlideService = Depends(get_slide_service),
    docx_generation_service: LegacyDocxGenerationService = Depends(
        get_legacy_docx_generation_service
    ),
    subscription_service: SubscriptionService = Depends(get_subscription_service)
):
    """
    [Synchronous] Upload DOCX -> Wait -> Return Slides JSON
    """
    if not file.filename.endswith('.docx'):
        raise HTTPException(400, "Chỉ chấp nhận tệp DOCX")

    # Get full user from repository with subscription data
    user_repo = SqliteUserRepository()
    user = user_repo.get_by_id(current_user.id)

    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    # Check subscription limits for docx upload
    subscription_service.check_feature_access(user, "docx_upload")

    # 1. Setup paths
    original_filename = Path(file.filename).name
    project_name = os.path.splitext(original_filename)[0]
    output_folder = prepare_output_folder(project_name)

    temp_path = _new_temp_docx_path(original_filename)
    with open(temp_path, "wb") as f:
        content = await file.read()
        f.write(content)

    try:
        # Estimate slide count and validate subscription limits. Keeping this
        # inside the try/finally also removes rejected uploads immediately.
        _check_docx_slide_limit(temp_path, subscription_service, user)

        # Increment usage counter
        subscription_service.increment_usage(current_user.id, "docx_upload")

        # 2. Run Workflow
        logger.info(f"🤖 Processing DOCX: {file.filename}")
        if not docx_generation_service.is_available():
            raise HTTPException(503, "Dịch vụ workflow không khả dụng")

        normalized_output_format = output_format.strip().lower()
        if normalized_output_format != "xml":
            raise HTTPException(400, "This workflow is XML-only; output_format must be xml")

        slide_options = {
            "output_format": normalized_output_format,
            "max_docx_slides": user.get_daily_limits()["max_docx_slides"],
        }
        output_folder_path = await docx_generation_service.generate_from_docx(
            docx_path=str(temp_path),
            output_folder=str(output_folder),
            slide_options=slide_options,
        )

        if not output_folder_path:
            raise HTTPException(
                status_code=500,
                detail=(
                    "Không tạo được slide từ tài liệu. "
                    "Vui lòng kiểm tra nội dung hoặc rút ngắn tài liệu rồi thử lại."
                ),
            )

        # 3. Import Project
        project_service.import_project_from_workflow(
            user_id=current_user.id,
            project_name=project_name,
            source_type='docx',
            source_data={'filename': original_filename},
            slide_options=slide_options,
            output_folder=output_folder
        )

        # 4. Build Response
        slides = slide_service.build_slides_from_output(str(output_folder), normalized_output_format)
        return JSONResponse({"slides": slides, "project_name": project_name})

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"DOCX Upload failed: {e}")
        error_payload = _workflow_error_payload(e)
        raise HTTPException(
            status_code=error_payload.get("status_code", 500),
            detail=error_payload["message"],
        )
    finally:
        _safe_unlink(temp_path)


@router.post("/docx/prepare-stream", response_model=DocxPrepareResponse)
async def prepare_docx_stream(
    file: UploadFile = File(...),
    workflow_mode: str = Form("auto_no_heading"),
    enable_transition_slides: str = Form("true"),
    enable_thank_you_slide: str = Form("true"),
    output_format: str = Form("xml"),
    current_user: UserResponse = Depends(get_current_active_user),
    subscription_service: SubscriptionService = Depends(get_subscription_service)
):
    """
    [Step 1 Stream] Upload DOCX & Save Options -> Return Filename
    """
    if not file.filename.endswith('.docx'):
        raise HTTPException(400, "Chỉ chấp nhận tệp DOCX")

    # Get full user from repository with subscription data
    user_repo = SqliteUserRepository()
    user = user_repo.get_by_id(current_user.id)

    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    # Check subscription limits for docx upload
    subscription_service.check_feature_access(user, "docx_upload")

    _cleanup_temp_dir()

    original_filename = Path(file.filename).name
    temp_path = _new_temp_docx_path(original_filename)
    content = await file.read()
    with open(temp_path, "wb") as f:
        f.write(content)

    try:
        _check_docx_slide_limit(temp_path, subscription_service, user)
    except Exception:
        _safe_unlink(temp_path)
        raise

    normalized_output_format = output_format.strip().lower()
    if normalized_output_format != "xml":
        raise HTTPException(400, "This workflow is XML-only; output_format must be xml")

    options = {
        "workflow_mode": workflow_mode.strip() or "auto_no_heading",
        "enable_transition_slides": enable_transition_slides.lower() in ('true', '1', 'yes', 'on'),
        "enable_thank_you_slide": enable_thank_you_slide.lower() in ('true', '1', 'yes', 'on'),
        "output_format": normalized_output_format,
        "max_docx_slides": user.get_daily_limits()["max_docx_slides"],
        "original_filename": original_filename,
    }
    stored_filename = temp_path.name
    with open(TEMP_DIR / f"{stored_filename}.options.json", "w") as f:
        json.dump(options, f)

    return {"filename": stored_filename}


# ==========================================
# MARKDOWN PARSER - Reusable for document preprocessing and Manual
# ==========================================

@router.post("/markdown/parse")
async def parse_markdown(
    content: dict,
    current_user: UserResponse = Depends(get_current_active_user)
):
    """
    Parse markdown text into hierarchical structure.
    Can be used by document preprocessing and Manual Mode.
    
    Request body:
    {
        "markdown": "# Title\\nContent...",
        "default_title": "Optional title"
    }
    """
    try:
        from backend.workflows.slidegen_workflow.utils.markdown_parser import markdown_to_manual_structure, validate_markdown_structure
        
        markdown_text = content.get('markdown', '').strip()
        default_title = content.get('default_title', 'Document')
        
        if not markdown_text:
            raise ValueError("Markdown content is empty")
        
        # Parse markdown
        document_title, structure = markdown_to_manual_structure(markdown_text, default_title)
        
        # Validate
        is_valid, errors = validate_markdown_structure(structure)
        if not is_valid:
            logger.warning(f"Markdown validation issues: {errors}")
        
        return {
            "success": True,
            "document_title": document_title,
            "structure": structure,
            "validation_warnings": errors if not is_valid else [],
            "node_count": len(structure)
        }
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        logger.error(f"Markdown parse error: {e}")
        raise HTTPException(status_code=500, detail="Failed to parse markdown")


@router.get("/docx/stream")
async def stream_docx_generation(
    filename: str = Query(..., description="DOCX filename"),
    current_user: UserResponse = Depends(get_current_user_from_token_param),
    project_service: ProjectService = Depends(get_project_service),
    docx_generation_service: LegacyDocxGenerationService = Depends(
        get_legacy_docx_generation_service
    ),
    subscription_service: SubscriptionService = Depends(get_subscription_service)
):
    """
    [Step 2 Stream] SSE Generation for DOCX
    """
    filename = os.path.basename(filename)
    temp_path = TEMP_DIR / filename
    if not temp_path.exists(): raise HTTPException(400, "Tệp không tồn tại hoặc đã hết hạn")

    # Increment usage counter (actual processing starts here)
    subscription_service.increment_usage(current_user.id, "docx_upload")

    # Load options
    options_path = TEMP_DIR / f"{filename}.options.json"
    slide_options = {}
    if options_path.exists():
        with open(options_path, "r") as f: slide_options = json.load(f)

    original_filename = Path(slide_options.get("original_filename", filename)).name
    project_name = os.path.splitext(original_filename)[0]
    output_folder = prepare_output_folder(project_name)

    async def event_generator():
        queue = asyncio.Queue()
        
        def callback(event_type, data):
            safe_data = sanitize_event_data(event_type, data)
            queue.put_nowait((event_type, safe_data))

        async def run_bg():
            res = await docx_generation_service.generate_from_docx(
                docx_path=str(temp_path),
                output_folder=str(output_folder),
                callback=callback,
                slide_options=slide_options
            )
            
            if not res:
                raise RuntimeError(
                    "Workflow không tạo được slide từ tài liệu. "
                    "Vui lòng rút ngắn tài liệu rồi thử lại."
                )

            pid = project_service.import_project_from_workflow(
                user_id=current_user.id,
                project_name=project_name,
                source_type='docx',
                source_data={'filename': original_filename},
                slide_options=slide_options,
                output_folder=output_folder
            )

            # Lấy số lượng slide thực tế từ kết quả workflow thay vì slide_options
            actual_slide_count = count_workflow_slides(
                output_folder,
                slide_options.get("output_format", "xml"),
            )

            workflow_temp_dir = str(output_folder)

            await queue.put((EventTypes.COMPLETE, {
                'project_id': pid,
                'project_name': project_name,
                'total_slides': actual_slide_count,
                'temp_dir': workflow_temp_dir
            }))

        asyncio.create_task(_run_stream_task(queue, run_bg, "DOCX"))

        # Common streaming loop
        while True:
            try:
                ev = await asyncio.wait_for(queue.get(), timeout=1.0)
                yield f"event: {ev[0]}\ndata: {json.dumps(ev[1])}\n\n"
                if ev[0] in [EventTypes.COMPLETE, EventTypes.ERROR_EVENT]: break
            except asyncio.TimeoutError:
                yield ": keepalive\n\n"
        
        # Cleanup input files after stream. Locked files are logged and ignored.
        _safe_unlink(temp_path)
        _safe_unlink(options_path)

    return StreamingResponse(event_generator(), media_type="text/event-stream")


@router.post("/docx/manual-session", response_model=ManualSessionResponse)
async def create_manual_session_from_docx(
    file: UploadFile = File(...),
    output_format: str = Form("xml"),
    current_user: UserResponse = Depends(get_current_active_user),
    structure_service: StructureService = Depends(get_structure_service),
    subscription_service: SubscriptionService = Depends(get_subscription_service)
):
    """
    Upload DOCX -> Parse Text -> Create Editing Session for Manual Mode
    """
    if not file.filename.endswith('.docx'):
        raise HTTPException(400, "Chỉ chấp nhận tệp DOCX")

    # Get full user from repository with subscription data
    user_repo = SqliteUserRepository()
    user = user_repo.get_by_id(current_user.id)

    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    # Check subscription limits for docx upload
    subscription_service.check_feature_access(user, "docx_upload")

    original_filename = Path(file.filename).name
    temp_path = _new_temp_docx_path(original_filename)
    content = await file.read()
    with open(temp_path, "wb") as f:
        f.write(content)

    try:
        _check_docx_slide_limit(temp_path, subscription_service, user)
    except Exception:
        _safe_unlink(temp_path)
        raise

    try:
        normalized_output_format = output_format.strip().lower()
        if normalized_output_format != "xml":
            raise HTTPException(400, "This workflow is XML-only; output_format must be xml")

        # Increment usage counter
        subscription_service.increment_usage(current_user.id, "docx_upload")

        from docx import Document
        doc = Document(temp_path)
        full_text = ""
        paragraphs_data = []
        for idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                start_pos = len(full_text)
                full_text += text + "\n\n"
                paragraphs_data.append({
                    "index": idx, "text": text,
                    "start_pos": start_pos, "end_pos": len(full_text)
                })

        session_id = structure_service.create_session(
            filename=original_filename,
            full_text=full_text,
            paragraphs=paragraphs_data,
            docx_path=str(temp_path),
            output_format=normalized_output_format,
        )

        return {
        "success": True,
        "session_id": session_id,
        "document_title": file.filename.replace('.docx', ''),
        "output_format": normalized_output_format,
        }
    except Exception as e:
        logger.error(f"Manual upload error: {e}")
        raise HTTPException(500, str(e))


# ==========================================
# 3. GENERATE FROM MANUAL STRUCTURE
# ==========================================


@router.get("/manual/stream")
async def stream_manual_generation(
    session_id: str,
    current_user: UserResponse = Depends(get_current_user_from_token_param),
    project_service: ProjectService = Depends(get_project_service),
    slide_workflow_service: SlideGenWorkflowService = Depends(
        get_slidegen_workflow_service
    )
):
    """Step 2 Manual: SSE Generation"""
    file_path = TEMP_DIR / f"approved_structure_{session_id}.json"
    if not file_path.exists(): raise HTTPException(404, "Không tìm thấy cấu trúc")
    
    with open(file_path, 'r', encoding='utf-8') as f: data = json.load(f)
    
    title = data.get('document_title', 'Untitled')
    output_folder = prepare_output_folder(title)

    async def event_generator():
        queue = asyncio.Queue()
        
        def callback(event_type, data):
            safe_data = sanitize_event_data(event_type, data)
            queue.put_nowait((event_type, safe_data))

        async def run_bg():
            slide_options = {
                'enable_transition_slides': data.get('enable_transition_slides', False),
                'enable_thank_you_slide': data.get('enable_thank_you_slide', False),
                'output_format': 'xml',
            }
            slide_chunks = data.get('slide_chunks', [])
            logger.info(f"[Manual Stream] Starting with {len(slide_chunks)} chunks, options={slide_options}")
            res = await slide_workflow_service.generate_from_chunks(
                slide_chunks=slide_chunks,
                title=title,
                output_folder=str(output_folder),
                callback=callback,
                slide_options=slide_options,
                slide_positions=data.get("slide_positions", []),
                slide_context_maps=data.get("slide_context_maps", []),
                section_ids=data.get("section_ids", []),
            )
            
            if not res:
                raise RuntimeError(
                    "Workflow không tạo được slide. Vui lòng thử lại."
                )

            pid = project_service.import_project_from_workflow(
                user_id=current_user.id,
                project_name=title,
                source_type='manual_structure',
                source_data={'session_id': session_id, 'title': title},
                slide_options={**slide_options, 'manual_mode': True},
                output_folder=output_folder
            )

            workflow_temp_dir = str(output_folder)

            # Derive actual generated slide count from XML files to avoid 0-total regressions.
            actual_slide_count = 0
            try:
                actual_slide_count = count_workflow_slides(
                    output_folder,
                    slide_options.get('output_format', 'xml'),
                )
            except Exception as count_err:
                logger.warning(f"[Manual Stream] Could not count generated slides: {count_err}")

            if actual_slide_count <= 0:
                actual_slide_count = len(data.get('slide_chunks', []))

            await queue.put((EventTypes.COMPLETE, {
                'project_id': pid,
                'project_name': title,
                'total_slides': actual_slide_count,
                'temp_dir': workflow_temp_dir
            }))

        asyncio.create_task(_run_stream_task(queue, run_bg, "MANUAL"))

        while True:
            try:
                ev = await asyncio.wait_for(queue.get(), timeout=1.0)
                yield f"event: {ev[0]}\ndata: {json.dumps(ev[1])}\n\n"
                if ev[0] in [EventTypes.COMPLETE, EventTypes.ERROR_EVENT]: break
            except asyncio.TimeoutError:
                yield ": keepalive\n\n"

    return StreamingResponse(event_generator(), media_type="text/event-stream")
