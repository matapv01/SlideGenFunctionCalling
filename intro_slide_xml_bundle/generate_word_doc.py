import os
import glob
from docx import Document

def main():
    folder_path = r"c:\Users\hp\Documents\SlideGenFunctionCalling_minh_rl\intro_slide_xml_bundle"
    xml_files = glob.glob(os.path.join(folder_path, "*.xml"))
    pptx_files = glob.glob(os.path.join(folder_path, "*.pptx"))
    
    # We just want unique slide concepts, let's extract base names
    slide_names = set()
    
    for f in xml_files + pptx_files:
        basename = os.path.basename(f)
        if basename.startswith("~") or basename == "xml_tool.py":
            continue
            
        # Remove extension
        name = os.path.splitext(basename)[0]
        
        # Remove suffixes like _test, _vX, _final, _premium
        name = name.replace("_test", "").replace("_final", "").replace("_premium", "")
        import re
        name = re.sub(r'_v\d+', '', name)
        
        # Replace underscores with spaces and title case
        name = name.replace("_", " ").title()
        
        if name and name.lower() not in ['convert universal', 'test slides', 'xml tool', 'tool des xml']:
            slide_names.add(name)

    doc = Document()
    doc.add_heading('Danh sách các slide', 0)
    
    for name in sorted(list(slide_names)):
        doc.add_paragraph(name, style='List Bullet')
        
    output_path = os.path.join(folder_path, "Danh_sach_slide.docx")
    doc.save(output_path)
    print(f"Created Word document at: {output_path}")

if __name__ == "__main__":
    main()
